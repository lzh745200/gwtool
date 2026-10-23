#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""跨平台真机验收脚本（麒麟 ARM64 / Windows / 通用 Linux）。

目的：让验收人员在本机**一条命令**跑出可回传的报告（零网络、零新依赖）：
  · 环境采集：OS/发行版/架构/Python/Qt/关键依赖版本、中文字体、屏幕与 DPI、
    数据目录与种子库；
  · 功能 smoke（全部在临时目录 + 临时库内完成，不碰用户真实库）：
      建库 → 导入 txt/md → 生成 docx → 内置渲染 PDF（量页数与正文行距）→
      备份后恢复 → 纠错 1 万字计时 → 写作提示 1 万字计时；
  · 输出 `accept_report_<时间戳>.json` 与同名 `.md`（人读摘要）。

用法（在仓库根目录或解包目录下）：
    python scripts/accept_env.py                 # 完整验收（含耗时项）
    python scripts/accept_env.py --fast          # 跳过耗时项（U 盘/低配机快速自检）
    python scripts/accept_env.py --json          # 只把 JSON 打到 stdout（不写文件）
    python scripts/accept_env.py --out /tmp/acc  # 报告写到指定目录

退出码：0 = 没有 FAIL（WARN/SKIP 允许）；1 = 有 FAIL；2 = 脚本自身无法启动
（参数/输出目录不可写等）。环境能力缺失（缺 Qt/PyMuPDF/中文字体等）只标 WARN，
不改变退出码，便于在最小安装的麒麟机器上也能先拿到环境报告。
"""
from __future__ import annotations

import argparse
import contextlib
import json
import os
import platform
import shutil
import statistics
import sys
import tempfile
import time
import traceback
from collections import Counter
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

REPORT_SCHEMA = 1
TOOL_NAME = "gwtool-accept-env"
PASS, WARN, FAIL, SKIP = "PASS", "WARN", "FAIL", "SKIP"
STATUSES = (PASS, WARN, FAIL, SKIP)

# 验收判据（与 doc/麒麟真机验收.md 的判据表保持一致）
EXPECTED_LINE_SPACING_PT = 28.0
LINE_SPACING_TOLERANCE_PT = 1.5
CORRECTOR_10K_LIMIT_MS = 2000.0
HINTS_10K_LIMIT_MS = 3000.0
MIN_PDF_PAGES = 2

SAMPLE_HEADING = "一、验收样张说明"
SAMPLE_BODY = (
    "验收正文样本段落：本段用于核对公文字体、字号、行距与分页是否正确。"
    "为深入贯彻落实上级关于进一步加强和改进新形势下安全生产工作的各项部署要求，"
    "请各部门结合实际细化措施、压实责任，确保各项任务落地见效。"
) * 12
SAMPLE_TAIL = "结语：以上为验收样张，仅用于环境核验，不包含真实公文内容。"
# 耗时项样本：含确定的常见错别字/易混词，保证"有命中"这个判据可判定
PERF_TEXT = ("工作布署已经完成，报名截止本月底，务必抓好落实。"
             * 500)[:10000]


@contextlib.contextmanager
def _quiet_native_stdout():
    """把**文件描述符 1 临时改道到 stderr**，兜住 C 扩展直接写 fd 的噪声。

    为什么 Python 层 `redirect_stdout` 不够：MuPDF（PyMuPDF 的 C 层）在遇到
    它不认识的 PDF 字典键时，会用 C 的 stdout 直接打印 `MuPDF error: ...`，
    完全绕过 `sys.stdout`。Linux CI 上这些行出现在 JSON 之前，导致
    `--fast --json` 的 stdout 不是纯 JSON（v1.11.0 首次 CI 实测）。
    这里在 **fd 层** dup2 重定向，并顺手关掉 PyMuPDF 的错误显示
    （不同版本 API 不同，一律 best-effort，失败不影响主流程）。
    """
    try:
        import pymupdf as _mupdf
        tools = getattr(_mupdf, "TOOLS", None)
        if tools is not None:
            for name in ("mupdf_display_errors", "mupdf_display_warnings"):
                fn = getattr(tools, name, None)
                if callable(fn):
                    try:
                        fn(False)
                    except Exception:      # 版本差异，忽略
                        pass
    except Exception:                      # 没装/旧版
        pass
    try:
        sys.stdout.flush()
        saved = os.dup(1)
    except (OSError, AttributeError, ValueError):
        yield                              # 拿不到 fd（如被重定向成 StringIO）就退化
        return
    try:
        os.dup2(2, 1)
        yield
    finally:
        try:
            sys.stdout.flush()
            os.dup2(saved, 1)
        finally:
            os.close(saved)


class Report:
    """收集检查项与指标，最后统一渲染 JSON/Markdown。"""

    def __init__(self) -> None:
        self.checks: list[dict] = []
        self.started = time.monotonic()
        self.environment: dict = {}
        self.artifacts: dict = {}

    def add(self, cid: str, title: str, status: str, detail: str = "",
            metrics: dict | None = None) -> None:
        assert status in STATUSES, status
        self.checks.append({"id": cid, "title": title, "status": status,
                            "detail": detail or "", "metrics": metrics or {}})

    def counts(self) -> dict:
        c = Counter(x["status"] for x in self.checks)
        fail = int(c.get(FAIL, 0))
        warn = int(c.get(WARN, 0))
        verdict = FAIL if fail else (WARN if warn else PASS)
        return {"pass": int(c.get(PASS, 0)), "warn": warn, "fail": fail,
                "skip": int(c.get(SKIP, 0)), "verdict": verdict}

    def payload(self, mode: str) -> dict:
        counts = self.counts()
        return {
            "schema": REPORT_SCHEMA,
            "tool": TOOL_NAME,
            "app_version": _app_version(),
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "mode": mode,
            "platform": self.environment.get("platform", {}),
            "environment": self.environment,
            "checks": self.checks,
            "summary": counts,
            "artifacts": self.artifacts,
            "duration_s": round(time.monotonic() - self.started, 2),
        }


def _app_version() -> str:
    try:
        from gwtool import __version__
        return str(__version__)
    except Exception:
        return "unknown"


def _dist_name() -> str:
    """发行版名称（Linux 读 os-release；其它平台用 platform.platform）。"""
    try:
        if sys.platform.startswith("linux"):
            p = Path("/etc/os-release")
            if p.exists():
                for line in p.read_text(encoding="utf-8",
                                        errors="replace").splitlines():
                    if line.startswith("PRETTY_NAME="):
                        return line.split("=", 1)[1].strip().strip('"')
    except Exception:
        pass
    return platform.platform()


def _dist_version(dist: str) -> str:
    try:
        from importlib import metadata
        return metadata.version(dist)
    except Exception:
        return "未安装"


def _product_data_dir_hint() -> Path:
    """产品默认数据目录路径（只算路径，不创建、不写探针）。"""
    if sys.platform.startswith("win"):
        base = os.environ.get("APPDATA") or str(
            Path.home() / "AppData" / "Roaming")
    else:
        base = os.environ.get("XDG_DATA_HOME") or str(
            Path.home() / ".local" / "share")
    return Path(base) / "gwtool"


def _ensure_gui():
    """创建 QGuiApplication（无显示环境自动退 offscreen），失败返回 None。"""
    if (not sys.platform.startswith("win")
            and not os.environ.get("DISPLAY")
            and not os.environ.get("WAYLAND_DISPLAY")):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    try:
        from PySide6.QtGui import QGuiApplication
    except Exception:
        return None
    try:
        return QGuiApplication.instance() or QGuiApplication([])
    except Exception:
        try:
            os.environ["QT_QPA_PLATFORM"] = "offscreen"
            return QGuiApplication.instance() or QGuiApplication([])
        except Exception:
            return None


CJK_FONT_HINTS = (
    "Noto Sans CJK", "Noto Serif CJK", "Noto Sans SC", "Noto Serif SC",
    "Source Han", "思源", "WenQuanYi", "文泉驿", "Droid Sans Fallback",
    "Microsoft YaHei", "微软雅黑", "SimSun", "宋体", "SimHei", "黑体",
    "KaiTi", "楷体", "FangSong", "仿宋", "AR PL", "Ubuntu", "DejaVu",
)


def _font_info(app) -> dict:
    info = {"families_count": 0, "cjk_fonts": [], "fallback_injected": ""}
    if app is None:
        return info
    try:
        from PySide6.QtGui import QFontDatabase
        fams = list(QFontDatabase.families())
        info["families_count"] = len(fams)
        info["cjk_fonts"] = sorted({f for f in fams
                                    if any(h.lower() in f.lower()
                                           for h in CJK_FONT_HINTS)})[:50]
    except Exception as exc:
        info["error"] = f"{type(exc).__name__}: {exc}"
        return info
    try:
        from gwtool.core.pdfrender import ensure_cjk_font
        injected = ensure_cjk_font()
        info["fallback_injected"] = str(injected or "")
    except Exception as exc:
        info["fallback_error"] = f"{type(exc).__name__}: {exc}"
    return info


def _screen_info(app) -> dict:
    if app is None:
        return {"available": False, "platform": os.environ.get(
            "QT_QPA_PLATFORM", "")}
    try:
        from PySide6.QtGui import QGuiApplication
        scr = QGuiApplication.primaryScreen()
        if scr is None:
            return {"available": False}
        geo = scr.geometry()
        avail = scr.availableGeometry()
        return {
            "available": True,
            "geometry": [geo.width(), geo.height()],
            "available_geometry": [avail.width(), avail.height()],
            "device_pixel_ratio": round(float(scr.devicePixelRatio()), 3),
            "logical_dots_per_inch": round(float(scr.logicalDotsPerInch()), 1),
            "qt_platform": os.environ.get("QT_QPA_PLATFORM", ""),
        }
    except Exception as exc:
        return {"available": False, "error": f"{type(exc).__name__}: {exc}"}


def collect_environment(app, report: Report) -> None:
    """采集环境信息，并把"能力是否存在"登记为 WARN/PASS 检查项。"""
    py_ver = platform.python_version()
    env = {
        "platform": {
            "os": platform.system(),
            "release": platform.release(),
            "version": platform.version(),
            "distribution": _dist_name(),
            "machine": platform.machine(),
            "processor": platform.processor(),
            "python": py_ver,
            "python_implementation": platform.python_implementation(),
            "python_executable": sys.executable,
            "frozen": bool(getattr(sys, "frozen", False)),
            "cwd": str(Path.cwd()),
        },
        "versions": {
            "gwtool": _app_version(),
            "PySide6": _dist_version("PySide6"),
            "PyMuPDF": _dist_version("PyMuPDF"),
            "jieba": _dist_version("jieba"),
            "python-docx": _dist_version("python-docx"),
            "numpy": _dist_version("numpy"),
            "lxml": _dist_version("lxml"),
            "sqlite": _sqlite_version(),
        },
        "fonts": _font_info(app),
        "screen": _screen_info(app),
        "paths": {
            "product_data_dir_hint": str(_product_data_dir_hint()),
            "seed_db": str(ROOT / "gwtool" / "resources" / "data" / "seed.db"),
            "seed_db_exists": (ROOT / "gwtool" / "resources" / "data"
                               / "seed.db").exists(),
        },
    }
    env["fonts"]["cjk_available"] = bool(
        env["fonts"]["cjk_fonts"] or env["fonts"].get("fallback_injected"))
    report.environment = env

    report.add("env.python", "Python 版本", PASS if _py_tuple(py_ver) >= (3, 9)
               else WARN, f"Python {py_ver}（目标 ≥3.9）")
    report.add("env.qt", "PySide6 / Qt", PASS if env["versions"]["PySide6"]
               != "未安装" else WARN,
               f"PySide6 {env['versions']['PySide6']}；screen="
               f"{'available' if env['screen'].get('available') else 'unavailable'}")
    report.add("env.pymupdf", "PyMuPDF", PASS if env["versions"]["PyMuPDF"]
               != "未安装" else WARN,
               f"PyMuPDF {env['versions']['PyMuPDF']}")
    report.add("env.jieba", "jieba 中文分词", PASS if env["versions"]["jieba"]
               != "未安装" else WARN,
               f"jieba {env['versions']['jieba']}")
    report.add("env.fonts", "中文字体", PASS if env["fonts"]["cjk_available"]
               else WARN,
               f"字体族 {env['fonts']['families_count']} 个；CJK 候选 "
               f"{len(env['fonts']['cjk_fonts'])} 个；兜底="
               f"{env['fonts'].get('fallback_injected') or '无'}")
    screen = env["screen"]
    screen_ok = bool(screen.get("available")) \
        and screen.get("qt_platform") != "offscreen"
    report.add("env.screen", "屏幕与 DPI",
               PASS if screen_ok else WARN,
               json.dumps(screen, ensure_ascii=False)
               + ("" if screen_ok else "（离屏/无真实显示会话）"))
    report.add("env.seed_db", "种子词典库", PASS if env["paths"]["seed_db_exists"]
               else WARN, env["paths"]["seed_db"])


def _py_tuple(ver: str) -> tuple:
    try:
        parts = [int(x) for x in ver.split(".")[:3]]
        return tuple(parts + [0] * (3 - len(parts)))
    except Exception:
        return (0, 0, 0)


def _sqlite_version() -> str:
    try:
        import sqlite3
        return sqlite3.sqlite_version
    except Exception:
        return "unknown"


def _imports_ok(module_names: list[str]) -> "tuple[bool, str]":
    missing = []
    for name in module_names:
        try:
            __import__(name)
        except Exception:
            missing.append(name)
    return (not missing), "、".join(missing)


def _write_sample(directory: Path, name: str, text: str) -> Path:
    p = directory / name
    p.write_text(text, encoding="utf-8")
    return p


def run_smoke(report: Report, workspace: Path, fast: bool) -> None:
    """功能 smoke：临时目录 + 临时库，不依赖 Office/COM，不触碰真实库。"""
    ok, missing = _imports_ok(["docx"])
    if not ok:
        report.add("smoke.import", "导入 txt/md", WARN,
                   f"缺少依赖：{missing}")
        report.add("smoke.docx", "生成 DOCX", WARN, f"缺少依赖：{missing}")
        report.add("smoke.pdf", "渲染 PDF", SKIP, "缺少 python-docx，跳过")
        report.add("smoke.backup_restore", "备份/恢复", SKIP, "缺少依赖")
        report.add("perf.corrector", "纠错 1 万字", SKIP, "--fast/缺依赖")
        report.add("perf.writing_hints", "写作提示 1 万字", SKIP, "--fast/缺依赖")
        return

    import docx  # 能力探测；后面用 docx.Document 回读生成物
    from gwtool import paths
    from gwtool.core import backup, corrector, docxgen, importer
    from gwtool.core import pdfrender, template, writing_hints
    from gwtool.db import connection as dbconn
    from gwtool.db import dao
    from gwtool.db.schema import init_schema

    smoke = workspace / "smoke"
    data_dir = smoke / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    paths.set_app_data_dir(data_dir)
    db_file = smoke / "accept.db"
    dbconn.configure(db_file)

    # ---------------------------------------------------------- 建库
    try:
        conn = dbconn.get_conn()
        init_schema(conn)
        ver = conn.execute("PRAGMA user_version").fetchone()[0]
        tables = conn.execute(
            "SELECT count(*) FROM sqlite_master WHERE type='table'").fetchone()[0]
        ok = bool(conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='documents'"
        ).fetchone())
        report.add("smoke.db", "建库 / 迁移", PASS if ok else FAIL,
                   f"user_version={ver}，表 {tables} 个")
    except Exception as exc:
        report.add("smoke.db", "建库 / 迁移", _exc_status(exc),
                   f"{type(exc).__name__}: {exc}")
        return

    # ---------------------------------------------------------- 导入 txt/md
    trees = []
    try:
        sample_dir = smoke / "samples"
        sample_dir.mkdir(parents=True, exist_ok=True)
        files = [
            _write_sample(sample_dir, "验收样例.txt",
                          SAMPLE_HEADING + "\n" + SAMPLE_BODY + "\n" + SAMPLE_TAIL),
            _write_sample(sample_dir, "验收样例.md",
                          f"# 验收样例（Markdown）\n\n{SAMPLE_HEADING}\n\n"
                          f"{SAMPLE_BODY}\n\n{SAMPLE_TAIL}\n"),
        ]
        blocked = []
        for f in files:
            r = importer.parse_any(str(f))
            if not r.ok or r.tree is None:
                blocked.append(f"{f.name}: {r.error or '解析失败'}")
                continue
            trees.append(r.tree)
            dao.add_document(dao.Document(
                title=r.tree.title or f.stem,
                content_text=r.tree.plain_text(),
                blocks_json=r.tree.to_json(),
                file_path=str(f), file_type=f.suffix.lstrip(".")))
        docs = dao.count_documents()
        ok = len(trees) == 2 and docs == 2
        report.add("smoke.import", "导入 txt/md", PASS if ok else FAIL,
                   f"解析 {len(trees)}/2 篇，入库 {docs} 篇"
                   + (f"；失败：{blocked}" if blocked else ""),
                   {"documents": docs})
    except Exception as exc:
        report.add("smoke.import", "导入 txt/md", _exc_status(exc),
                   f"{type(exc).__name__}: {exc}")
        return

    # ---------------------------------------------------------- 生成 DOCX
    sample_tree = _build_sample_tree()
    docx_path = smoke / "accept_sample.docx"
    try:
        docxgen.generate_docx([sample_tree] + trees, template.default_template(),
                              str(docx_path))
        d = docx.Document(str(docx_path))
        text = "\n".join(p.text for p in d.paragraphs)
        ok = docx_path.exists() and docx_path.stat().st_size > 3000 \
            and "验收正文样本段落" in text
        report.add("smoke.docx", "生成 DOCX", PASS if ok else FAIL,
                   f"{docx_path.name} {docx_path.stat().st_size // 1024} KB，"
                   f"段落 {len(d.paragraphs)} 个",
                   {"bytes": docx_path.stat().st_size})
    except Exception as exc:
        report.add("smoke.docx", "生成 DOCX", _exc_status(exc),
                   f"{type(exc).__name__}: {exc}")

    # ---------------------------------------------------------- 渲染 PDF + 行距
    pdf_path = smoke / "accept_sample.pdf"
    try:
        tpl = template.default_template()
        pdfrender.render_compiled_pdf([sample_tree], tpl, str(pdf_path))
        metrics = _measure_pdf(pdf_path, tpl)
        spacing = metrics.get("body_line_spacing_pt")
        pages = int(metrics.get("pages") or 0)
        status = PASS
        detail = (f"页数 {pages}；正文行距 "
                  f"{'—' if spacing is None else f'{spacing:.2f}pt'}；"
                  f"正文汉字 {metrics.get('chinese_chars', 0)} 个")
        if pages < MIN_PDF_PAGES:
            status = FAIL
            detail += f"（应 ≥{MIN_PDF_PAGES} 页）"
        if spacing is None:
            status = WARN if status == PASS else status
            detail += "（无法测出行距，可能未取到正文页）"
        elif abs(spacing - EXPECTED_LINE_SPACING_PT) > LINE_SPACING_TOLERANCE_PT:
            status = FAIL
            detail += (f"（判据 {EXPECTED_LINE_SPACING_PT}±"
                       f"{LINE_SPACING_TOLERANCE_PT}pt）")
        if metrics.get("chinese_chars", 0) == 0:
            status = WARN if status == PASS else status
            detail += "（PDF 未提取到中文，字体可能未生效）"
        report.add("smoke.pdf", "渲染 PDF / 行距", status, detail, metrics)
    except Exception as exc:
        report.add("smoke.pdf", "渲染 PDF / 行距", _exc_status(exc),
                   f"{type(exc).__name__}: {exc}")

    # ---------------------------------------------------------- 备份 / 恢复
    try:
        n_before = dao.count_documents()
        t0 = time.perf_counter()
        rep = backup.create_backup_detailed(note="accept-env",
                                            mode=backup.MODE_MANUAL)
        backup_s = time.perf_counter() - t0
        size = Path(rep.path).stat().st_size
        dao.add_document(dao.Document(title="恢复前新增", content_text="备份后新增",
                                      blocks_json="[]", file_type="txt"))
        n_mid = dao.count_documents()
        t0 = time.perf_counter()
        r = backup.restore_backup_detailed(str(rep.path))
        restore_s = time.perf_counter() - t0
        n_after = dao.count_documents()
        ok = r.ok and n_after == n_before and n_mid == n_before + 1
        report.add("smoke.backup_restore", "备份 / 恢复", PASS if ok else FAIL,
                   f"备份 {size // 1024} KB / {backup_s:.2f}s；恢复 "
                   f"{restore_s:.2f}s；文档 {n_mid}→{n_after}（备份时 {n_before}）",
                   {"backup_bytes": size, "backup_s": round(backup_s, 3),
                    "restore_s": round(restore_s, 3),
                    "documents_before": n_before, "documents_after": n_after})
    except Exception as exc:
        report.add("smoke.backup_restore", "备份 / 恢复", _exc_status(exc),
                   f"{type(exc).__name__}: {exc}")

    # ---------------------------------------------------------- 纠错计时
    perf_text = PERF_TEXT
    if fast:
        report.add("perf.corrector", "纠错 1 万字", SKIP,
                   "--fast：跳过耗时项", {"chars": len(perf_text)})
        report.add("perf.writing_hints", "写作提示 1 万字", SKIP,
                   "--fast：跳过耗时项", {"chars": len(perf_text)})
        return
    try:
        t0 = time.perf_counter()
        hits = corrector.check_text(perf_text)
        ms = (time.perf_counter() - t0) * 1000
        ok = ms <= CORRECTOR_10K_LIMIT_MS and bool(hits)
        report.add("perf.corrector", "纠错 1 万字", PASS if ok else FAIL,
                   f"{len(perf_text)} 字 / {ms:.0f} ms / 命中 {len(hits)} 处"
                   f"（上限 {CORRECTOR_10K_LIMIT_MS:.0f} ms）",
                   {"chars": len(perf_text), "ms": round(ms, 1),
                    "hits": len(hits)})
    except Exception as exc:
        report.add("perf.corrector", "纠错 1 万字", _exc_status(exc),
                   f"{type(exc).__name__}: {exc}")

    try:
        hint_text = PERF_TEXT
        t0 = time.perf_counter()
        hs = writing_hints.hints(hint_text, "通知")
        ms = (time.perf_counter() - t0) * 1000
        ok = ms <= HINTS_10K_LIMIT_MS and bool(hs)
        report.add("perf.writing_hints", "写作提示 1 万字", PASS if ok else FAIL,
                   f"{len(hint_text)} 字 / {ms:.0f} ms / 提示 {len(hs)} 条"
                   f"（上限 {HINTS_10K_LIMIT_MS:.0f} ms）",
                   {"chars": len(hint_text), "ms": round(ms, 1),
                    "hints": len(hs)})
    except Exception as exc:
        report.add("perf.writing_hints", "写作提示 1 万字", _exc_status(exc),
                   f"{type(exc).__name__}: {exc}")


def _build_sample_tree():
    from gwtool.core.model import Block, DocTree, HEADING, PARAGRAPH
    return DocTree(title="验收样张", blocks=[
        Block(type=HEADING, level=1, text="验收样张"),
        Block(type=HEADING, level=2, text=SAMPLE_HEADING),
        Block(type=PARAGRAPH, text=SAMPLE_BODY),
        Block(type=PARAGRAPH, text=SAMPLE_BODY),
        Block(type=PARAGRAPH, text=SAMPLE_TAIL),
    ])


def _measure_pdf(pdf_path: Path, tpl) -> dict:
    """用 PyMuPDF 量页数、正文行距与中文字符数（不依赖 Office）。"""
    # 优先新包名：`import fitz` 在 PyMuPDF 1.24+ 会往 stdout 打弃用提示，
    # 会污染 --json 的纯 JSON 输出。
    try:
        import pymupdf as fitz
    except Exception:                 # 老版本只有 fitz
        import fitz  # type: ignore

    if not Path(pdf_path).exists():
        raise FileNotFoundError(f"PDF 未生成：{pdf_path}")
    d = fitz.open(str(pdf_path))
    try:
        pages = d.page_count
        m = {"pages": pages, "body_page": -1, "body_line_spacing_pt": None,
             "line_samples": 0, "chinese_chars": 0}
        mm = 72.0 / 25.4
        body_page = None
        for i in range(pages):
            try:
                if d[i].search_for("验收正文样本段落"):
                    body_page = i
                    break
            except Exception:
                continue
        if body_page is None:
            return m
        m["body_page"] = body_page
        page = d[body_page]
        top = max(0.0, float(tpl.margin_top_mm) * mm - 8)
        bottom = page.rect.height - max(0.0, float(tpl.margin_bottom_mm) * mm - 8)
        body_size = float(tpl.body_size_pt)
        deltas: list[float] = []
        all_ys: list[float] = []
        try:
            data = page.get_text("dict")
        except Exception:
            data = {}
        for block in data.get("blocks", []):
            if block.get("type") != 0:
                continue
            lines = []
            for ln in block.get("lines", []):
                spans = [sp for sp in ln.get("spans", []) if sp.get("text", "").strip()]
                if not spans:
                    continue
                size = max(float(sp.get("size", 0)) for sp in spans)
                if body_size * 0.8 <= size <= body_size * 1.2:
                    y = float(spans[0].get("origin", (0, 0))[1])
                    if top <= y <= bottom:
                        lines.append(y)
            lines.sort()
            all_ys.extend(lines)
            for a, b in zip(lines, lines[1:]):
                delta = b - a
                if 10.0 <= delta <= 45.0:
                    deltas.append(delta)
        if not deltas and len(all_ys) >= 2:
            # Qt 的 PDF 写出器常把每个文本行拆成独立 block（同一段落的行
            # 不共享 block），按整页正文行统一排序再求相邻行距。
            all_ys.sort()
            for a, b in zip(all_ys, all_ys[1:]):
                delta = b - a
                if 20.0 <= delta <= 36.0:
                    deltas.append(delta)
        if deltas:
            m["body_line_spacing_pt"] = round(statistics.median(deltas), 2)
            m["line_samples"] = len(deltas)
        text = page.get_text()
        m["chinese_chars"] = sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")
        if tpl.page_number_enabled:
            last = d[pages - 1].get_text()
            m["page_number_stamped"] = "—" in last and str(pages) in last
        return m
    finally:
        d.close()


def _exc_status(exc: BaseException) -> str:
    """缺依赖/无显示等环境问题记 WARN；其它异常是真失败。"""
    if isinstance(exc, (ImportError, ModuleNotFoundError)):
        return WARN
    if "could not connect to display" in str(exc).lower():
        return WARN
    return FAIL


def _write_markdown(report: dict, out_path: Path) -> None:
    c = report["summary"]
    lines = [
        "# 公文汇编助手 · 环境验收报告",
        "",
        f"- 版本：{report['app_version']}",
        f"- 生成时间：{report['generated_at']}",
        f"- 模式：{'快速（--fast，跳过耗时项）' if report['mode'] == 'fast' else '完整'}",
        f"- 结论：**{c['verdict']}**（PASS {c['pass']} / WARN {c['warn']} / "
        f"FAIL {c['fail']} / SKIP {c['skip']}）",
        "",
        "## 1. 环境",
        "",
        "| 项 | 值 |",
        "|---|---|",
    ]
    pf = report.get("platform", {})
    env = report.get("environment", {})
    for key, label in (("os", "操作系统"), ("distribution", "发行版"),
                       ("machine", "架构"), ("python", "Python"),
                       ("python_implementation", "实现")):
        lines.append(f"| {label} | {pf.get(key, '')} |")
    for key, label in (("PySide6", "PySide6"), ("PyMuPDF", "PyMuPDF"),
                       ("jieba", "jieba"), ("python-docx", "python-docx"),
                       ("sqlite", "SQLite")):
        lines.append(f"| {label} | {env.get('versions', {}).get(key, '')} |")
    fonts = env.get("fonts", {})
    lines.append(f"| 中文字体 | 字体族 {fonts.get('families_count', 0)} 个；"
                 f"CJK 候选 {len(fonts.get('cjk_fonts') or [])} 个；"
                 f"兜底 {fonts.get('fallback_injected') or '无'} |")
    scr = env.get("screen", {})
    lines.append(f"| 屏幕 | {json.dumps(scr, ensure_ascii=False)} |")
    paths = env.get("paths", {})
    for key, label in (("product_data_dir_hint", "默认数据目录"),
                       ("seed_db", "种子库")):
        lines.append(f"| {label} | {paths.get(key, '')} |")
    lines += ["", "## 2. 检查项", "", "| 检查 | 状态 | 说明 |", "|---|---|---|"]
    for item in report["checks"]:
        lines.append(f"| {item['title']} | **{item['status']}** | "
                     f"{item['detail']} |")
    lines += [
        "",
        "## 3. 判据",
        "",
        f"- PDF 正文行距：{EXPECTED_LINE_SPACING_PT} ± "
        f"{LINE_SPACING_TOLERANCE_PT} pt；样张页数 ≥ {MIN_PDF_PAGES}",
        f"- 纠错 1 万字：≤ {CORRECTOR_10K_LIMIT_MS:.0f} ms 且有命中",
        f"- 写作提示 1 万字：≤ {HINTS_10K_LIMIT_MS:.0f} ms 且有结果",
        "- 备份/恢复：恢复后文档数与备份时点一致、恢复过程无异常",
        "",
        "## 4. 回传",
        "",
        "把本报告的同名 `.json` 与 `.md` 一起回传（JSON 是机器判据，"
        "Markdown 是人读摘要）。报告生成在：",
        "",
    ]
    for name, path in (report.get("artifacts") or {}).items():
        lines.append(f"- {name}: `{path}`")
    lines += [
        "",
        "> 本报告不含任何真实公文正文：功能 smoke 全部使用临时目录与内置样例。",
        "",
    ]
    out_path.write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        prog="accept_env.py",
        description="跨平台真机验收（环境采集 + 功能 smoke，零网络）")
    ap.add_argument("--out", default="", help="报告输出目录（默认产品数据目录）")
    ap.add_argument("--fast", action="store_true", help="跳过耗时项（纠错/提示计时）")
    ap.add_argument("--json", action="store_true",
                    help="只把 JSON 打印到 stdout（不写报告文件）")
    args = ap.parse_args(argv)

    report = Report()
    workspace = Path(tempfile.mkdtemp(prefix="gwtool_accept_"))
    app = None
    try:
        # `--json` 要求 stdout 是纯 JSON：把整段"干活"过程放进 fd 级静音区，
        # C 扩展（MuPDF/Qt）直接写 fd 1 的噪声会落到 stderr。
        quiet = _quiet_native_stdout() if args.json else contextlib.nullcontext()
        with quiet:
            app = _ensure_gui()
            collect_environment(app, report)

            out_dir: Path | None = None
            if not args.json:
                out_dir = Path(args.out) if args.out else _product_data_dir_hint()
                try:
                    out_dir.mkdir(parents=True, exist_ok=True)
                    probe = out_dir / ".gwtool_accept_write_test"
                    probe.write_text("", encoding="utf-8")
                    probe.unlink()
                except OSError as exc:
                    print(f"报告目录不可写：{out_dir}（{exc}）", file=sys.stderr)
                    return 2

            run_smoke(report, workspace, fast=args.fast)

        payload = report.payload("fast" if args.fast else "full")
        if args.json:
            print(json.dumps(payload, ensure_ascii=False, indent=1))
            return 0 if payload["summary"]["fail"] == 0 else 1

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        json_path = out_dir / f"accept_report_{stamp}.json"
        md_path = out_dir / f"accept_report_{stamp}.md"
        payload["artifacts"] = {"json": str(json_path), "markdown": str(md_path)}
        json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=1),
                             encoding="utf-8")
        _write_markdown(payload, md_path)
        c = payload["summary"]
        print(f"验收结论：{c['verdict']}（PASS {c['pass']} / WARN {c['warn']} / "
              f"FAIL {c['fail']} / SKIP {c['skip']}）")
        for item in payload["checks"]:
            print(f"  [{item['status']:4}] {item['title']}：{item['detail']}")
        print(f"报告已写入：\n  {json_path}\n  {md_path}")
        return 0 if c["fail"] == 0 else 1
    except Exception as exc:
        print(f"验收脚本异常：{type(exc).__name__}: {exc}", file=sys.stderr)
        traceback.print_exc(file=sys.stderr)
        return 2
    finally:
        try:
            from gwtool.db import connection as _dbconn
            _dbconn.close_current_thread()
        except Exception:
            pass
        try:
            from gwtool import paths as _paths
            _paths.set_app_data_dir(None)
        except Exception:
            pass
        shutil.rmtree(workspace, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
