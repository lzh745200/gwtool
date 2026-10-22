# -*- coding: utf-8 -*-
"""校对报告导出（DOCX 留痕）测试。

分三层守：
  · core/report 的报告树与 DOCX 产物 —— 明细行数 == 命中数、空命中写
    “未发现问题”、提示类必须标注“提示类，建议人工确认”、修正后正文不丢不重；
  · AnyDocCorrectDialog 的入口 —— 按钮存在、导出成功/失败只改 lbl_stat、
    导出中按钮置灰、没有命中时也允许导出（留一份“已校对”凭证）；
  · 布局护栏 —— row2 加到第 5 个按钮后，1080 宽窗口仍放得下。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document as DX

from gwtool import __version__
from gwtool.core import corrector, report
from gwtool.core.model import TABLE

SAMPLE_TEXT = "关于布署工作的通知\n报名截止8月底，会议大约三天左右召开。"


# ------------------------------------------------------------ 工具
def _text_of(path: Path) -> str:
    """DOCX 全部段落 + 表格单元格文本（拼成一段便于断言）。"""
    d = DX(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _detail_rows(path: Path) -> list[list[str]]:
    """报告里的明细表行（去掉表头）；没有明细表则返回空列表。"""
    d = DX(str(path))
    for t in d.tables:
        rows = [[c.text for c in row.cells] for row in t.rows]
        if rows and rows[0][:1] == ["序号"]:
            return rows[1:]
    return []


def _export(tmp_path: Path, hits, **kwargs) -> Path:
    out = tmp_path / "校对报告.docx"
    report.export_correction_report(hits, str(out), **kwargs)
    assert out.exists() and out.stat().st_size > 2000, "报告 DOCX 未生成"
    return out


def _make_dialog(qapp, text: str, source: str = "测试公文"):
    """构造对话框并直接喂入命中。

    不走 `_load_text`/`start_check`：那会真的起扫描线程，`ok` 是队列连接，
    等它排到事件循环时可能在本用例后续 `wait_bg` 里再触发一次 `_on_checked`，
    把“已导出”状态覆盖掉。这里用与后台线程同一个 `_scan_blocks` 同步算，
    `_blocks`/`_source` 直接设置，测试只关心报告导出这一条链路。
    """
    from gwtool.ui.correct_dialog import AnyDocCorrectDialog, _scan_blocks

    dlg = AnyDocCorrectDialog()
    dlg._blocks = [{"kind": "para", "text": line, "rows": None, "level": 0}
                   for line in text.splitlines()]
    dlg._source = source
    dlg._on_checked(_scan_blocks(dlg._blocks))
    return dlg


# ------------------------------------------------------------ 报告树
def test_build_tree_has_required_metadata():
    hits = corrector.check_text(SAMPLE_TEXT)
    tree = report.build_correction_report_tree(hits, source_name="示例公文.docx")
    text = "\n".join(b.text for b in tree.blocks if b.text)

    assert tree.title == "公文校对报告"
    assert "校对对象：示例公文.docx" in text
    assert "校对时间：" in text
    assert "《党政机关公文格式》GB/T 9704" in text
    assert __version__ in text, "校对依据必须带版本说明"
    assert "校对结论：" in text and "未发现问题" not in text

    table = next(b for b in tree.blocks if b.type == TABLE)
    assert table.rows[0] == ["序号", "类别", "原文", "建议", "置信度", "依据"]
    assert len(table.rows) - 1 == len(hits)
    # 只报“第几处”，不报块内偏移（多块时偏移没有全局意义）
    assert [r[0] for r in table.rows[1:]] == [str(i) for i in range(1, len(hits) + 1)]


def test_export_report_detail_rows_match_hits(tmp_path):
    hits = corrector.check_text(SAMPLE_TEXT)
    assert len(hits) >= 2
    out = _export(tmp_path, hits, source_name="示例公文.docx")
    rows = _detail_rows(out)
    assert len(rows) == len(hits), "明细行数必须等于命中数"
    for row, c in zip(rows, hits):
        assert row[1] == c.category
        assert row[2] == c.wrong
        assert row[4] == f"{c.confidence:.2f}"

    text = _text_of(out)
    for key in ("校对对象：示例公文.docx", "校对依据：", "校对结论：",
                "GB/T 9704", __version__, "确认错误", "疑似", "提示类"):
        assert key in text, f"报告缺少字段：{key}"


def test_export_empty_hits_says_no_problem(tmp_path):
    out = _export(tmp_path, [], source_name="无问题公文.docx")
    text = _text_of(out)
    assert "未发现问题" in text
    assert "校对对象：无问题公文.docx" in text
    assert _detail_rows(out) == [], "空命中不应生成明细表"


def test_nested_hits_from_multiple_blocks(tmp_path):
    """UI 侧 `_corrs` 是 list[list[Correction]]，报告入口必须能直接吃。"""
    groups = [corrector.check_text("工作布署已经完成。"),
              corrector.check_text("会议大约三天左右召开。")]
    out = _export(tmp_path, groups)
    assert len(_detail_rows(out)) == sum(len(g) for g in groups)


def test_export_fixed_text_keeps_content_once(tmp_path):
    hits = corrector.check_text("工作布署已经完成。")
    fixed = "工作部署已经完成。\n第二段内容保持不变。"
    out = _export(tmp_path, hits, source_name="公文.docx", fixed_text=fixed)
    text = _text_of(out)
    assert "修正后正文" in text
    assert text.count("修正后正文") == 1, "修正后正文节不得重复出现"
    assert text.count("工作部署已经完成。") == 1, "修正后正文不丢内容也不重复"
    assert "第二段内容保持不变。" in text


def test_advisory_hits_flagged_manual_confirm(tmp_path):
    hits = corrector.check_text("会议大约三天左右召开，报名截止8月底。")
    advisory = [c for c in hits if c.category in corrector.ADVISORY_CATEGORIES]
    normal = [c for c in hits if c.category not in corrector.ADVISORY_CATEGORIES]
    assert advisory, "样例应至少有提示类命中"
    assert normal, "样例应至少有普通命中"

    out = _export(tmp_path, hits)
    rows = _detail_rows(out)
    lucky = [r for r in rows if r[1] in corrector.ADVISORY_CATEGORIES]
    assert len(lucky) == len(advisory)
    # 提示类的 suggestion 是标签不是替换文本：必须显式标注，不能让人误以为
    # 可以直接按“建议”列替换原文。
    assert all("提示类，建议人工确认" in r[3] for r in lucky)
    plain = [r for r in rows if r[1] not in corrector.ADVISORY_CATEGORIES]
    assert plain and all("提示类" not in r[3] for r in plain)
    assert "提示类，建议人工确认" in _text_of(out)


# ------------------------------------------------------------ UI 入口
def test_dialog_export_report_success(tmp_db, qapp, monkeypatch, wait_bg, tmp_path):
    from gwtool.ui import correct_dialog as cd

    dlg = _make_dialog(qapp, SAMPLE_TEXT, source="测试公文.docx")
    try:
        assert dlg.btn_report.text() == "导出报告"
        assert dlg.btn_report.isEnabled()
        n = dlg.listw.count()
        assert n == len(dlg._hits) >= 2

        out = tmp_path / "对话框校对报告.docx"
        monkeypatch.setattr(
            cd.QFileDialog, "getSaveFileName",
            staticmethod(lambda *a, **k: (str(out), "")))
        dlg.export_report()
        assert not dlg.btn_report.isEnabled(), "导出期间「导出报告」应置灰"
        assert wait_bg(dlg._export_worker, cond=out.exists), "导出任务未在限期内完成"
        assert out.exists()
        assert "已导出校对报告" in dlg.lbl_stat.text()
        assert dlg.btn_report.isEnabled(), "导出完成后按钮应复位"

        assert len(_detail_rows(out)) == n, "明细行数应等于对话框里的命中数"
        text = _text_of(out)
        assert "校对对象：测试公文.docx" in text
        assert "提示类，建议人工确认" in text
        assert "修正后正文" in text and "关于布署工作的通知" in text
    finally:
        dlg.deleteLater()


def test_dialog_export_without_hits_allowed(tmp_db, qapp, monkeypatch, wait_bg, tmp_path):
    from gwtool.ui import correct_dialog as cd

    dlg = _make_dialog(qapp, "今天天气很好。", source="干净公文.docx")
    try:
        assert dlg.listw.count() == 0
        assert dlg.btn_report.isEnabled(), "没有命中时也应允许导出（结论写“未发现问题”）"

        out = tmp_path / "无命中校对报告.docx"
        monkeypatch.setattr(
            cd.QFileDialog, "getSaveFileName",
            staticmethod(lambda *a, **k: (str(out), "")))
        dlg.export_report()
        assert wait_bg(dlg._export_worker, cond=out.exists)
        assert "未发现问题" in _text_of(out)
        assert _detail_rows(out) == []
    finally:
        dlg.deleteLater()


def test_dialog_report_export_failure_only_updates_status(
        tmp_db, qapp, monkeypatch, wait_bg, tmp_path):
    from gwtool.core import report as report_mod
    from gwtool.ui import correct_dialog as cd

    dlg = _make_dialog(qapp, SAMPLE_TEXT)
    try:
        def boom(*a, **k):
            raise OSError("磁盘已满")

        monkeypatch.setattr(report_mod, "export_correction_report", boom)
        out = tmp_path / "失败.docx"
        monkeypatch.setattr(
            cd.QFileDialog, "getSaveFileName",
            staticmethod(lambda *a, **k: (str(out), "")))
        dlg.export_report()
        assert wait_bg(dlg._export_worker)
        assert "导出失败" in dlg.lbl_stat.text()
        assert not out.exists()
        assert dlg.btn_report.isEnabled(), "失败后按钮必须复位"
    finally:
        dlg.deleteLater()


def test_dialog_report_button_follows_busy_state(tmp_db, qapp):
    from gwtool.ui.correct_dialog import AnyDocCorrectDialog

    dlg = AnyDocCorrectDialog()
    try:
        assert not dlg.btn_report.isEnabled(), "未载入内容时导出按钮应置灰"
        dlg._blocks = [{"kind": "para", "text": "正文", "rows": None, "level": 0}]
        dlg._set_busy(True)
        for name in ("btn_docx", "btn_txt", "btn_report"):
            assert not getattr(dlg, name).isEnabled(), name
        dlg._set_busy(False)
        assert dlg.btn_report.isEnabled(), "扫描结束后导出按钮应恢复"
    finally:
        dlg.deleteLater()


def test_dialog_row2_fits_default_window(qapp):
    from gwtool.ui.correct_dialog import AnyDocCorrectDialog

    dlg = AnyDocCorrectDialog()
    try:
        assert dlg.width() == 1080
        # row2 从 4 个按钮加到 5 个后，最小宽度仍必须落在窗口内 ——
        # 若按钮撑破布局，minimumSizeHint 会超过 1080（窗口会被迫变宽）。
        assert dlg.minimumSizeHint().width() <= 1080, "第 5 个按钮把最小宽度顶出了窗口"
        for name in ("btn_prev", "btn_next", "btn_docx", "btn_txt", "btn_report"):
            assert getattr(dlg, name).text()
        assert dlg.btn_report.text() == "导出报告"
    finally:
        dlg.deleteLater()
