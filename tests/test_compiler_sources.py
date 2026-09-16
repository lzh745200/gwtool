# -*- coding: utf-8 -*-
"""汇编材料来源清单（C3）的测试。

汇编类公文需要**可追溯性**：领导或审计问"这份汇编里的内容哪来的"，
必须答得出来。此前成品是一个孤立 DOCX，来源只存在于操作者记忆里。

同时要守住：默认不改变既有行为（`include_sources` 默认 False），
否则所有老调用方产出的文件会凭空多出一段附录。
"""
from __future__ import annotations

from pathlib import Path

from gwtool.core import compiler
from gwtool.db import dao


def _add(title: str, **kw) -> int:
    # 注意：不要试图指定 import_time —— DAO 在写入时自动填充，
    # 调用方传的值会被覆盖（测试里断言具体时间会失败）。
    base = dict(title=title, content_text=f"{title}的正文内容",
                blocks_json="[]", file_path=str(Path("/tmp") / f"{title}.docx"))
    base.update(kw)
    return dao.add_document(dao.Document(**base))


class TestCollectSources:
    def test_collects_library_documents(self, tmp_db):
        did = _add("材料甲")
        got = compiler.collect_sources([did], [])
        assert len(got) == 1
        assert got[0]["title"] == "材料甲"
        assert got[0]["file"] == "材料甲.docx"
        assert got[0]["time"], "导入时间应由 DAO 自动填充"

    def test_marks_extra_paths_as_not_imported(self, tmp_db, tmp_path):
        p = tmp_path / "外部材料.docx"
        p.write_bytes(b"x")
        got = compiler.collect_sources([], [str(p)])
        assert got[0]["category"] == "（未入库）"
        assert got[0]["file"] == "外部材料.docx"

    def test_skips_deleted_documents(self, tmp_db):
        """材料在汇编前被删掉了就跳过，不该编造一行出来。"""
        did = _add("已删除材料")
        dao.delete_document(did)
        assert compiler.collect_sources([did], []) == []

    def test_skips_missing_ids(self, tmp_db):
        assert compiler.collect_sources([999999], []) == []

    def test_order_matches_input(self, tmp_db):
        a = _add("甲")
        b = _add("乙")
        got = compiler.collect_sources([b, a], [])
        assert [g["title"] for g in got] == ["乙", "甲"]

    def test_survives_category_lookup_failure(self, tmp_db, monkeypatch):
        """分类查不到不能让整条来源收集失败。"""
        did = _add("材料甲")

        def boom():
            raise RuntimeError("分类表炸了")
        monkeypatch.setattr(dao, "list_categories", boom)
        got = compiler.collect_sources([did], [])
        assert len(got) == 1


class TestBuildTree:
    def test_structure(self):
        tree = compiler.build_sources_tree(
            [{"title": "甲", "file": "a.docx", "time": "2026-09-01 10:00",
              "category": "综合"}])
        assert tree.title == compiler.SOURCES_TITLE
        table = [b for b in tree.blocks if b.type == "table"]
        assert len(table) == 1
        rows = table[0].rows
        assert rows[0] == ["序号", "材料标题", "原文件名", "导入时间", "所属分类"]
        assert rows[1] == ["1", "甲", "a.docx", "2026-09-01 10:00", "综合"]

    def test_numbering_starts_at_one(self):
        tree = compiler.build_sources_tree(
            [{"title": "甲"}, {"title": "乙"}])
        table = [b for b in tree.blocks if b.type == "table"][0]
        assert [r[0] for r in table.rows[1:]] == ["1", "2"]


class TestCompileIntegration:
    def _read_text(self, path) -> str:
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)

    def test_default_does_not_add_sources(self, tmp_db, tmp_path):
        """默认必须保持既有行为，否则所有老调用方的产物都会多出附录。"""
        did = _add("材料甲")
        out = tmp_path / "out.docx"
        compiler.compile_docx(compiler.CompileRequest(
            doc_ids=[did], out_docx=str(out)))
        assert compiler.SOURCES_TITLE not in self._read_text(out)

    def test_include_sources_appends_table(self, tmp_db, tmp_path):
        did = _add("材料甲")
        out = tmp_path / "out.docx"
        compiler.compile_docx(compiler.CompileRequest(
            doc_ids=[did], out_docx=str(out), include_sources=True))
        text = self._read_text(out)
        assert compiler.SOURCES_TITLE in text
        assert "材料甲.docx" in text

    def test_extra_paths_listed(self, tmp_db, tmp_path):
        from docx import Document
        src = tmp_path / "源材料.docx"
        d = Document()
        d.add_paragraph("外部材料正文")
        d.save(str(src))
        did = _add("材料甲")
        out = tmp_path / "out.docx"
        compiler.compile_docx(compiler.CompileRequest(
            doc_ids=[did], extra_paths=[str(src)], out_docx=str(out),
            include_sources=True))
        text = self._read_text(out)
        assert "源材料.docx" in text
        assert "（未入库）" in text

    def test_no_empty_appendix_when_no_materials(self, tmp_db, tmp_path):
        """没有材料时不追加只有表头的空附录。"""
        out = tmp_path / "out.docx"
        try:
            compiler.compile_docx(compiler.CompileRequest(
                doc_ids=[], extra_paths=[], out_docx=str(out),
                include_sources=True))
        except ValueError:
            return          # "没有可汇编的材料" 是既有行为，符合预期
        assert compiler.SOURCES_TITLE not in self._read_text(out)


class TestWiring:
    def test_worker_accepts_flag(self):
        import pathlib
        root = pathlib.Path(__file__).resolve().parent.parent
        src = (root / "gwtool" / "ui" / "workers.py").read_text(encoding="utf-8")
        assert "include_sources" in src

    def test_wizard_has_checkbox_default_on(self):
        import pathlib
        root = pathlib.Path(__file__).resolve().parent.parent
        src = (root / "gwtool" / "ui" / "compile_wizard.py").read_text(
            encoding="utf-8")
        assert "chk_sources" in src
        idx = src.index("self.chk_sources.setChecked(")
        assert "True" in src[idx:idx + 45], "来源清单默认应勾选"
        assert "include_sources=self.chk_sources.isChecked()" in src
