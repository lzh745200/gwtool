# -*- coding: utf-8 -*-
"""第十轮全面排查（v1.20.0 基线）的缺陷护栏。

每条用例都对应一个**修复前会失败**的真实行为，不是"看起来该这样"：

1. `paths._ensure_dir` 的可写探针用固定文件名 → 并发调用时互相删对方的探针，
   Windows 报 WinError 32 被翻译成「目录不可写」，启动直接拒绝（误报）。
2. `attachments.safe_stored_name` 不处理 Windows 保留设备名 → `CON.txt`
   这类附件「复制成功」但 0 字节、磁盘上没有文件，库里却登记成功。
3. `app.run` 在 QApplication 构造**之前**问系统深浅色 → 非 Windows 平台
   恒取到 Unknown，麒麟上「跟随系统深浅色」从未生效。
4. `exporter.build` 直接写用户选定路径 → 打包中断会在交付路径留下半截包。
5. `reminder.enabled` 读库失败返回 False，与自身文档「默认开启」矛盾。
"""
from __future__ import annotations

import inspect
import os
import threading
import zipfile

import pytest

from gwtool import paths
from gwtool.core import attachments, batch, backup, exporter, naming, reminder
from gwtool.db import dao


def _add(title: str, tmp_path, **kw) -> int:
    base = dict(title=title, content_text=f"{title}的正文", blocks_json="[]",
                file_path=str(tmp_path / "不存在" / f"{title}.docx"))
    base.update(kw)
    return dao.add_document(dao.Document(**base))


# ------------------------------------------------------------------ 1 路径探针
class TestEnsureDirConcurrency:
    def test_concurrent_calls_never_report_unwritable(self, tmp_path):
        """并发（多线程/多实例）调用不得把可写目录判成不可写。

        修复前：探针文件名固定为 `.gwtool_write_test`，两个调用方互相
        `unlink()` 对方的文件 → FileNotFoundError / WinError 32 →
        DataDirError「目录不可写」。实测 8 线程 × 200 次误报 1223 次。
        """
        target = tmp_path / "data" / "logs"
        errors: list = []
        lock = threading.Lock()

        def worker():
            for _ in range(50):
                try:
                    paths._ensure_dir(target)
                except paths.DataDirError as exc:      # 修复前会大量命中
                    with lock:
                        errors.append(str(exc))

        threads = [threading.Thread(target=worker) for _ in range(8)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

        assert errors == [], f"并发下误报目录不可写 {len(errors)} 次：{errors[:2]}"
        assert [p.name for p in target.iterdir()
                if "write_test" in p.name] == []


class TestEnsureDirProbeCaching:
    def test_probe_runs_only_once_per_directory(self, tmp_path, monkeypatch):
        """同一目录只做一次可写探针（否则逐条目循环里按条数重复付账）。

        实测：真实数据目录单次探针 ~45 ms，`attachments_dir()` 因两层探针
        89.7 ms/次；导出 300 个附件就会白等二十多秒。
        """
        target = tmp_path / "cached"
        paths._ensure_dir(target)           # 首次：真探测

        calls = {"n": 0}

        def counting(self, *a, **k):
            calls["n"] += 1
            raise OSError(28, "boom")

        monkeypatch.setattr(paths.Path, "write_text", counting)
        for _ in range(5):
            paths._ensure_dir(target)       # 已缓存：不应再探测
        assert calls["n"] == 0

    def test_cached_directory_is_recreated_when_removed(self, tmp_path):
        """缓存不能让"目录已被删"蒙混过关：快路径必须回退到重建。"""
        target = tmp_path / "gone"
        paths._ensure_dir(target)
        target.rmdir()
        assert paths._ensure_dir(target).is_dir()


# ------------------------------------------------------------------ 2 设备名
class TestWindowsDeviceNames:
    @pytest.mark.parametrize("name,expected", [
        ("CON.txt", "_CON.txt"),
        ("nul.pdf", "_nul.pdf"),
        ("COM1.docx", "_COM1.docx"),
        ("LPT9", "_LPT9"),
        ("普通.pdf", "普通.pdf"),
        ("报告.", "报告"),            # 尾点会被系统静默截掉
    ])
    def test_safe_stored_name(self, name, expected):
        assert attachments.safe_stored_name(name) == expected

    def test_both_call_sites_agree(self):
        """附件与批量汇编对同一名字必须给同一结论（防两份实现再次漂移）。"""
        for name in ("CON.txt", "con.docx", "PRN", "COM3.xlsx", "LPT1"):
            assert batch.safe_filename(name).startswith("_")
            assert attachments.safe_stored_name(name).startswith("_")

    def test_fix_windows_traps_is_idempotent(self):
        once = naming.fix_windows_traps("CON.txt")
        assert naming.fix_windows_traps(once) == once

    @pytest.mark.skipif(os.name != "nt", reason="Windows 保留设备名语义")
    def test_restore_writes_device_named_attachment(self, tmp_path, tmp_db):
        """端到端：备份包内的 `attachments/CON.txt` 必须真的落成文件。

        修复前实测：恢复计数 = 1（报告成功），附件目录**为空**；而
        `Path.exists()` 因设备语义返回 True，所以 `_missing_attachments`
        也不会把它报成缺失 —— 灾难恢复场景下静默丢件。
        """
        pkg = tmp_path / "pkg.zip"
        with zipfile.ZipFile(pkg, "w") as zf:
            zf.writestr("attachments/CON.txt", "附件正文".encode("utf-8"))

        with zipfile.ZipFile(pkg) as zf:
            n = backup._restore_attachments(zf)

        assert n == 1
        real = [p for p in paths.attachments_dir().iterdir() if p.is_file()]
        assert [p.name for p in real] == ["_CON.txt"]      # 已改名，内容完整
        assert real[0].read_text(encoding="utf-8") == "附件正文"


# ------------------------------------------------------------------ 3 深色主题
class TestDarkThemeOrdering:
    def test_system_dark_query_happens_after_qapplication(self):
        """`_system_prefers_dark` 必须排在 QApplication 构造之后。

        非 Windows 平台只有存在应用实例时 `styleHints().colorScheme()` 才给出
        真实值；排在构造前会恒得 Unknown，麒麟上「跟随系统深浅色」静默失效。
        """
        src = inspect.getsource(__import__(
            "gwtool.app", fromlist=["run"]).run)
        i_app = src.index("QApplication.instance()")
        i_dark = src.index("_system_prefers_dark()")
        assert i_app < i_dark, "系统深浅色判定被挪回 QApplication 构造之前"


# ------------------------------------------------------------------ 4 交付包
class TestExportAtomicity:
    def test_failed_export_leaves_no_partial_package(self, tmp_path, tmp_db,
                                                     monkeypatch):
        """打包中断时，用户选定路径上不能留下半截 zip，临时文件也要清掉。"""
        _add("甲", tmp_path)
        out = tmp_path / "pkg.zip"

        real_writestr = zipfile.ZipFile.writestr
        seen = {"n": 0}

        def boom(self, name, data, *a, **kw):
            seen["n"] += 1
            if seen["n"] >= 2:
                raise OSError(28, "No space left on device")
            return real_writestr(self, name, data, *a, **kw)

        monkeypatch.setattr(zipfile.ZipFile, "writestr", boom)
        with pytest.raises(OSError):
            exporter.build(exporter.ExportRequest(out_path=str(out)))

        assert not out.exists(), "中断后仍在交付路径留下文件"
        assert not (tmp_path / "pkg.zip.part").exists(), "临时文件未清理"

    def test_successful_export_replaces_into_place(self, tmp_path, tmp_db):
        """正常路径：.part 改名后包可用、且无残留。"""
        _add("甲", tmp_path)
        out = tmp_path / "pkg.zip"
        rep = exporter.build(exporter.ExportRequest(out_path=str(out)))
        assert rep["documents"] == 1
        assert out.exists()
        assert exporter.verify_package(out)["ok"] is True
        assert not (tmp_path / "pkg.zip.part").exists()


# ------------------------------------------------------------------ 5 提醒默认值
class TestReminderDefault:
    def test_enabled_defaults_true_when_setting_unreadable(self, monkeypatch):
        """读设置失败是环境问题，不能被解释成「用户关掉了提醒」。"""

        def boom(*a, **k):
            raise RuntimeError("database is locked")

        monkeypatch.setattr(reminder.dao, "get_setting", boom)
        assert reminder.enabled() is True


# ------------------------------------------------------------------ 6 纯文本重复标题
class TestPlainTextTitleDedup:
    """解析器把首个标题块兼作 title，`plain_text()` 不能再拼一遍。

    content_text 同时用于字数统计、查重哈希、FTS 索引与"导出纯文本"，
    重复一次会让短文档字数接近翻倍、导出件首行重影（修复前实测三种格式均重复）。
    """

    @pytest.mark.parametrize("name,body", [
        ("a.txt", "关于XX的通知\n各处室：\n正文。\n"),
        ("a.md", "# 关于XX的通知\n\n各处室：\n正文。\n"),
        ("b.html",
         "<html><body><h1>关于XX的通知</h1><p>各处室：</p><p>正文。</p></body></html>"),
    ])
    def test_imported_title_appears_once(self, tmp_path, name, body):
        from gwtool.core.importer import parse_any

        p = tmp_path / name
        p.write_text(body, encoding="utf-8")
        tree = parse_any(str(p)).tree
        assert tree.title == "关于XX的通知"
        assert tree.plain_text().count(tree.title) == 1

    def test_docx_uses_first_heading_as_title_without_duplicating(self, tmp_path):
        from docx import Document

        from gwtool.core.importer import parse_any

        p = tmp_path / "h.docx"
        doc = Document()
        doc.add_heading("关于XX的通知", level=1)
        doc.add_paragraph("各处室：")
        doc.save(p)

        tree = parse_any(str(p)).tree
        assert tree.title == "关于XX的通知"
        assert tree.plain_text().count(tree.title) == 1
