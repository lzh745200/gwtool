# -*- coding: utf-8 -*-
"""文档生成链（compiler/docxgen/watermark/booklet/template/skeletons/report）
逐模块深探：全部真实生成 docx/pdf 产物并读回验证，禁 mock。

此前零执行路径：
  - compiler.compile_docx 全链（空材料拒绝、extra 文件、旧数据兜底、标题覆盖）；
  - docxgen._build_cover（封面各字段真实渲染）；
  - docx_to_pdf 真机 COM/LibreOffice（无组件环境自动跳过）；
  - watermark 的单章模式/原地覆盖两条保存路径；
  - booklet 空页数/空 PDF/同路径另存边界；
  - template.from_json 非法配置、中文数字边界、未知键过滤；
  - skeletons title_hint 回退、report.verdict 纯提示分支。
"""
from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from gwtool.core import booklet
from gwtool.core import compiler
from gwtool.core import docxgen
from gwtool.core import report
from gwtool.core import skeletons
from gwtool.core import template as tplmod
from gwtool.core import watermark
from gwtool.core.model import Block, DocTree, PARAGRAPH
from gwtool.core.template import CoverInfo, DocTemplate
from gwtool.db import dao


def _tree(title: str, *paras: str) -> DocTree:
    t = DocTree(title=title)
    t.blocks = [Block(type=PARAGRAPH, text=p) for p in paras]
    return t


def _store_doc(title: str, content: str) -> int:
    tree = _tree(title, *content.splitlines())
    return dao.add_document(dao.Document(
        title=title, content_text=content, blocks_json=tree.to_json(),
        file_type="docx"))


# ---------------------------------------------------------------- compiler
class TestCompileDocxFullChain:
    def test_empty_request_rejected(self, tmp_db):
        with pytest.raises(ValueError, match="没有可汇编"):
            compiler.compile_docx(compiler.CompileRequest())

    def test_full_compile_with_extra_file(self, tmp_db, tmp_path):
        d1 = _store_doc("调研报告", "第一段内容\n第二段内容")
        d2 = _store_doc("会议纪要", "纪要正文")
        extra = tmp_path / "外部材料.txt"
        extra.write_text("外部补充材料正文", encoding="utf-8")
        out = tmp_path / "汇编成果.docx"
        req = compiler.CompileRequest(
            doc_ids=[d1, 999999, d2],          # 999999 不存在，应被跳过
            extra_paths=[str(extra)],
            out_docx=str(out))
        result = compiler.compile_docx(req)
        assert Path(result).exists()
        from docx import Document
        doc = Document(result)
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "调研报告" in all_text
        assert "会议纪要" in all_text
        assert "外部补充材料正文" in all_text

    def test_legacy_plain_text_fallback(self, tmp_db, tmp_path):
        """blocks_json 为空的旧数据：按纯文本逐段兜底，不产生空材料。"""
        did = dao.add_document(dao.Document(
            title="旧数据文档", content_text="旧第一行\n旧第二行",
            blocks_json="[]", file_type="txt"))
        out = tmp_path / "旧数据汇编.docx"
        compiler.compile_docx(compiler.CompileRequest(
            doc_ids=[did], out_docx=str(out)))
        assert Path(out).exists()

    def test_material_titles_override(self, tmp_db, tmp_path):
        did = _store_doc("原标题", "正文内容")
        out = tmp_path / "覆盖标题.docx"
        compiler.compile_docx(compiler.CompileRequest(
            doc_ids=[did], out_docx=str(out),
            material_titles=["新标题甲"]))
        from docx import Document
        all_text = "\n".join(p.text for p in Document(out).paragraphs)
        assert "新标题甲" in all_text
        assert "原标题" not in all_text


class TestDocxToPdfReal:
    def test_real_conversion_when_component_available(self, tmp_db, tmp_path):
        """本机有 WPS/Word/LibreOffice 时真实转换；都没有则跳过。

        有组件就必须真跑——此前该函数 0 覆盖，COM 分支从未执行过。
        """
        has_soffice = bool(shutil.which("soffice") or shutil.which("libreoffice"))
        com_ok = False
        if not has_soffice and shutil.os.name == "nt":
            # pywin32 可导入 ≠ 机器上真的有 WPS/Word（CI runner 就是这种）。
            # 用真 Dispatch 探测：任一组件能创建并正常退出才算可用。
            try:
                import win32com.client
                for prog_id in ("KWPS.Application", "Word.Application"):
                    try:
                        app = win32com.client.Dispatch(prog_id)
                        app.Quit()
                        com_ok = True
                        break
                    except Exception:
                        com_ok = False
            except ImportError:
                com_ok = False
        if not has_soffice and not com_ok:
            pytest.skip("本机无可用的 docx->PDF 组件")
        src = tmp_path / "待转.docx"
        trees = [_tree("转换样张", "用于转换的正文")]
        docxgen.generate_docx(trees, DocTemplate(), str(src))
        out = compiler.docx_to_pdf(str(src), str(tmp_path / "转出.pdf"))
        assert Path(out).exists() and Path(out).stat().st_size > 0


# ---------------------------------------------------------------- docxgen 封面
class TestCoverBuild:
    def test_cover_fields_rendered(self, tmp_db, tmp_path):
        tpl = DocTemplate()
        tpl.cover = CoverInfo(
            enabled=True, title="年度工作汇编", subtitle="内部资料",
            org="示例机关办公室", date="二〇二六年九月",
            extra_lines=["编印说明一", "编印说明二"])
        out = tmp_path / "带封面.docx"
        docxgen.generate_docx([_tree("汇编", "正文")], tpl, str(out))
        from docx import Document
        all_text = "\n".join(p.text for p in Document(out).paragraphs)
        for expected in ("年度工作汇编", "内部资料", "示例机关办公室",
                         "二〇二六年九月", "编印说明一", "编印说明二"):
            assert expected in all_text, f"封面缺 {expected}"


# ---------------------------------------------------------------- watermark
def _mk_pdf(path: Path, pages: int = 2) -> None:
    import pymupdf as fitz
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 100), f"第{i + 1}页测试内容", fontname="china-s",
                         fontsize=14)
    doc.save(str(path))
    doc.close()


class TestWatermarkPaths:
    def test_single_stamp_no_tile(self, tmp_path):
        src = tmp_path / "a.pdf"
        _mk_pdf(src)
        out = tmp_path / "a_stamped.pdf"
        watermark.stamp_watermark_pdf(str(src), "征求意见稿",
                                      out_path=str(out), tile=False)
        import pymupdf as fitz
        with fitz.open(str(out)) as doc:
            assert doc.page_count == 2
            text = doc[0].get_text()
        assert "征求意见稿" in text

    def test_pdf_overwrite_in_place(self, tmp_path):
        """默认覆盖原文件：走 临时文件 + os.replace 原子替换路径。"""
        src = tmp_path / "b.pdf"
        _mk_pdf(src, pages=1)
        before = src.stat().st_mtime_ns
        watermark.stamp_watermark_pdf(str(src), "秘密★1年")
        assert src.exists(), "原地覆盖后文件必须仍在"
        assert not Path(str(src) + ".wm.tmp").exists(), "临时文件必须已被替换掉"
        import pymupdf as fitz
        with fitz.open(str(src)) as doc:
            assert "秘密★1年" in doc[0].get_text()
        assert src.stat().st_mtime_ns >= before

    def test_docx_watermark_in_place(self, tmp_path):
        docx = tmp_path / "c.docx"
        docxgen.generate_docx([_tree("水印样张", "正文")], DocTemplate(), str(docx))
        out = watermark.add_watermark_docx(str(docx), "内部资料")
        assert out == str(docx), "未传 out_path 时原地保存"
        from docx import Document
        xml = Document(str(docx)).sections[0].header._element.xml
        assert "内部资料" in xml and "PowerPlusWaterMarkObject" in xml


# ---------------------------------------------------------------- booklet
class TestBookletEdges:
    def test_zero_or_negative_pages(self):
        assert booklet.booklet_order(0) == []
        assert booklet.booklet_order(-3) == []

    def test_empty_pdf_rejected(self, tmp_path):
        """零页 PDF（外部工具/损坏产物）：make_booklet 必须明确拒绝。

        pymupdf 自身不允许保存零页文档，因此用手工最小 PDF 字节构造。
        """
        src = tmp_path / "empty.pdf"
        src.write_bytes(
            b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
            b"2 0 obj<</Type/Pages/Kids[]/Count 0>>endobj\n"
            b"trailer<</Root 1 0 R/Size 3>>\n%%EOF")
        import pymupdf as fitz
        assert fitz.open(str(src)).page_count == 0  # 确为零页
        with pytest.raises(ValueError, match="源 PDF 为空"):
            booklet.make_booklet(str(src), str(tmp_path / "out.pdf"))

    def test_export_a4_same_path_no_rewrite(self, tmp_path):
        """同路径另存：不落盘重写，只返回页数（防御误覆盖）。"""
        src = tmp_path / "same.pdf"
        _mk_pdf(src, pages=3)
        before = src.read_bytes()
        n = booklet.export_a4_pdf(str(src), str(src))
        assert n == 3
        assert src.read_bytes() == before, "同路径不得重写文件"


# ---------------------------------------------------------------- template
class TestTemplateEdges:
    def test_from_json_non_dict_rejected(self):
        with pytest.raises(ValueError, match="非法"):
            DocTemplate.from_json('["不是对象"]')

    @pytest.mark.parametrize("n,expect", [
        (1, "一"), (9, "九"), (10, "十"), (11, "十一"),
        (19, "十九"), (20, "二十"), (21, "二十一"),
        (99, "九十九"), (100, "100"), (0, "0"),
    ])
    def test_cn_num_boundaries(self, n, expect):
        assert tplmod._cn_num(n) == expect

    def test_material_label(self):
        assert tplmod.material_label("", 3) == ""
        assert tplmod.material_label("材料：", 3) == "材料："
        assert tplmod.material_label("材料{n}：", 2) == "材料二："

    def test_known_fields_filters_unknown_keys(self):
        got = tplmod._known_fields(CoverInfo,
                                   {"title": "甲", "幽灵键": 1, "enabled": True})
        assert got == {"title": "甲", "enabled": True}
        assert tplmod._known_fields(CoverInfo, "不是字典") == {}


# ---------------------------------------------------------------- skeletons/report
class TestSkeletonHintFallback:
    def test_title_hint_used_when_no_title_kw(self):
        """不传 title 时用 title_hint 模板填充（占位符按 kw 替换）。"""
        sk = skeletons.get("通知")
        text = sk.render(org="示例机关", matter="全市安全生产")
        assert "示例机关" in text and "全市安全生产" in text
        assert "{org}" not in text and "{matter}" not in text

    def test_title_kw_wins_over_hint(self):
        sk = skeletons.get("通知")
        text = sk.render(title="自定义标题", org="示例机关", matter="事项")
        assert "自定义标题" in text


class TestReportVerdict:
    def test_verdict_branches(self):
        from gwtool.core.inspector import Finding
        assert report.verdict([]) == "未发现格式问题"
        info = [Finding(severity="info", item="检查项", detail="提示内容")]
        assert "提示" in report.verdict(info)
