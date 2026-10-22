# -*- coding: utf-8 -*-
"""v1.9 A 批排版保真回归（P1-01 / P1-03 / P1-23）。

三件事都是用**真实生成物**验的，不 mock 生成链路：
  P1-01  docx 段落间距：正文/标题段必须显式 before=0 after=0 line=560 exact，
         且 styles.xml 的 docDefaults 段前段后也清零（根因：python-docx 默认
         模板带 w:after=200，实测被 WPS 叠加成 32.4 磅行距、每面 19~20 行）。
  P1-03  PDF 转换前更新域：COM 分支必须先 Fields.Update()+Repaginate() 再
         SaveAs2，否则 WPS 直接另存得到的 PDF 目录页只有占位文字；
         LibreOffice 分支若产出"目录未更新"的 PDF，不得当作成功直接返回。
  P1-23  主送机关行双冒号：模板 "{recipients}：" + 用户输入自带冒号 -> "：："。

COM 部分用**假 Dispatch**（注入 sys.modules）驱动真实 `compiler.docx_to_pdf`，
不依赖本机装没装 WPS/Word，也不会真的拉起办公进程。
"""
from __future__ import annotations

import sys
import types
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from gwtool.core import compiler, docxgen, skeletons
from gwtool.core.model import Block, DocTree
from gwtool.core.template import DocTemplate, default_template

TOC_PLACEHOLDER = compiler._TOC_PLACEHOLDER


# ---------------------------------------------------------------- 公共样张
def _trees() -> list[DocTree]:
    t = DocTree(title="关于加强公文处理工作的通知")
    t.blocks = [
        Block(type="heading", level=1, text="一、总体要求"),
        Block(type="paragraph", text="各单位要高度重视公文处理工作，认真抓好落实。"),
        Block(type="heading", level=2, text="（一）提高思想认识"),
        Block(type="paragraph", text="充分认识做好此项工作的重要意义。"),
    ]
    return [t]


def _gen(tmp_path: Path, name: str = "out.docx", tpl=None) -> Path:
    out = tmp_path / name
    docxgen.generate_docx(_trees(), tpl or default_template(), str(out))
    assert out.exists() and out.stat().st_size > 0
    return out


def _spacing_attrs(par) -> dict:
    """段落 pPr/w:spacing 的属性字典（无该元素返回空字典）。"""
    ppr = par._p.pPr
    if ppr is None:
        return {}
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        return {}
    return {k.rsplit("}", 1)[-1]: v for k, v in spacing.attrib.items()}


def _doc_defaults_spacing(docx_path: Path) -> dict:
    """styles.xml -> docDefaults/pPrDefault/pPr/w:spacing 的属性字典。"""
    with zipfile.ZipFile(docx_path) as zf:
        xml = zf.read("word/styles.xml").decode("utf-8")
    from lxml import etree
    root = etree.fromstring(xml.encode("utf-8"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    node = root.find("w:docDefaults/w:pPrDefault/w:pPr/w:spacing", namespaces=ns)
    if node is None:
        return {}
    return {k.rsplit("}", 1)[-1]: v for k, v in node.attrib.items()}


# ================================================================ P1-01 行距
class TestParagraphSpacing:
    def test_body_paragraph_spacing_is_explicit_zero(self, tmp_path):
        """正文段：before/after=0 且 line=560、lineRule=exact（28 磅固定行距）。"""
        docx_path = _gen(tmp_path)
        doc = Document(str(docx_path))
        body = next(p for p in doc.paragraphs if p.text.startswith("各单位要高度重视"))
        attrs = _spacing_attrs(body)
        assert attrs.get("before") == "0", f"段前距未清零：{attrs}"
        assert attrs.get("after") == "0", f"段后距未清零：{attrs}"
        assert attrs.get("line") == "560", f"行距应为 28 磅(560 twips)：{attrs}"
        assert attrs.get("lineRule") == "exact", f"行距规则应为固定值：{attrs}"
        # 与模板参数一致（python-docx 侧读数）
        assert body.paragraph_format.line_spacing.pt == 28
        assert body.paragraph_format.space_after.pt == 0
        assert body.paragraph_format.space_before.pt == 0

    def test_heading_paragraph_spacing_is_explicit_zero(self, tmp_path):
        """标题段同样清零（Heading 样式自带 before/after，不能靠继承）。"""
        docx_path = _gen(tmp_path)
        doc = Document(str(docx_path))
        head = next(p for p in doc.paragraphs if p.text == "一、总体要求")
        attrs = _spacing_attrs(head)
        assert attrs.get("before") == "0"
        assert attrs.get("after") == "0"
        assert attrs.get("line") == "560"
        assert attrs.get("lineRule") == "exact"

    def test_doc_defaults_spacing_zeroed(self, tmp_path):
        """根因位置：styles.xml docDefaults 的段前段后必须清零。

        这是 P1-01 的关键一条 —— python-docx 默认模板带
        `<w:spacing w:after="200" w:line="276" w:lineRule="auto"/>`，
        只改段落不动它，表格单元格/分页段等仍会继承 10 磅段后距。
        """
        docx_path = _gen(tmp_path)
        attrs = _doc_defaults_spacing(docx_path)
        assert attrs, "docDefaults 里应存在 w:spacing（被清成 0 而不是删除）"
        assert attrs.get("before") == "0"
        assert attrs.get("after") == "0"
        # 不动 line/lineRule：未显式设行距的段落（表格等）保持模板默认 1.15 倍
        assert attrs.get("line") == "276"
        assert attrs.get("lineRule") == "auto"

    def test_custom_line_spacing_still_honored(self, tmp_path):
        """模板改行距时按参数写 XML（回归：不要让清零逻辑覆盖 line）。"""
        tpl = DocTemplate()
        tpl.line_spacing_pt = 20
        docx_path = _gen(tmp_path, "ls20.docx", tpl)
        doc = Document(str(docx_path))
        body = next(p for p in doc.paragraphs if p.text.startswith("各单位要高度重视"))
        attrs = _spacing_attrs(body)
        assert attrs.get("line") == "400"      # 20pt * 20 twips
        assert attrs.get("lineRule") == "exact"
        assert attrs.get("after") == "0"

    def test_table_cells_do_not_inherit_after_spacing(self, tmp_path):
        """表格单元格段落不得继承 docDefaults 的 10 磅段后距。"""
        t = DocTree(title="表格材料")
        t.blocks = [Block(type="table", rows=[["序号", "内容"], ["1", "测试"]])]
        out = tmp_path / "table.docx"
        docxgen.generate_docx([t], default_template(), str(out))
        doc = Document(str(out))
        cell = doc.tables[0].cell(1, 1)
        attrs = _spacing_attrs(cell.paragraphs[0])
        # 单元格没有显式 spacing 时，docDefaults 已清零 → 继承值即 0
        assert attrs.get("after", "0") in ("0", None)
        assert _doc_defaults_spacing(out).get("after") == "0"


# ================================================================ P1-03 域更新
def _write_pdf(path: Path, text: str = "") -> None:
    """写一份真 PDF（可带一段中文文字），供 _pdf_has_stale_toc 判定。"""
    import pymupdf as fitz
    doc = fitz.open()
    try:
        page = doc.new_page()
        if text:
            page.insert_text((72, 100), text, fontname="china-s", fontsize=12)
        doc.save(str(path))
    finally:
        doc.close()


def _docx_stub(tmp_path: Path) -> Path:
    """真实生成一份带目录域 + 占位串的 docx（不依赖 TOC 参数细节）。"""
    tpl = default_template()
    tpl.toc_enabled = True
    return _gen(tmp_path, "toc.docx", tpl)


class _FakeDoc:
    def __init__(self, log: list, pdf_path: Path, fail_update: bool = False):
        self._log = log
        self._pdf_path = pdf_path
        self._fail_update = fail_update
        self.Fields = types.SimpleNamespace(Update=self._update)

    def _update(self):
        self._log.append("Fields.Update")
        if self._fail_update:
            raise RuntimeError("WPS 拒绝更新域（模拟）")

    def Repaginate(self):
        self._log.append("Repaginate")

    def SaveAs2(self, path, FileFormat=17):
        self._log.append("SaveAs2")
        _write_pdf(Path(path), "")          # 不带占位串 = 目录已更新

    def Close(self, save=False):
        self._log.append("Close")


class _FakeApp:
    def __init__(self, log: list, pdf_path: Path, fail_update: bool = False):
        self._log = log
        self._pdf_path = pdf_path
        self._fail_update = fail_update
        self.Visible = False
        self.Documents = types.SimpleNamespace(Open=self._open)

    def _open(self, path, ReadOnly=True):
        self._log.append("Open")
        return _FakeDoc(self._log, self._pdf_path, self._fail_update)

    def Quit(self):
        self._log.append("Quit")


def _fake_win32com(log: list, pdf_path: Path, fail_update: bool = False):
    """构造一个可注入 sys.modules 的 win32com.client 替身。"""
    client = types.ModuleType("win32com.client")

    def _dispatch(progid):                  # 只记录，不拉起真实办公进程
        log.append(f"Dispatch:{progid}")
        return _FakeApp(log, pdf_path, fail_update)

    client.Dispatch = _dispatch
    pkg = types.ModuleType("win32com")
    pkg.client = client
    return pkg, client


def _install_fake_com(monkeypatch, log, pdf_path, fail_update=False):
    pkg, client = _fake_win32com(log, pdf_path, fail_update)
    monkeypatch.setitem(sys.modules, "win32com", pkg)
    monkeypatch.setitem(sys.modules, "win32com.client", client)


def _block_com(monkeypatch):
    """让 `import win32com.client` 必然 ImportError（模拟无 Office 的机器）。"""
    monkeypatch.setitem(sys.modules, "win32com", types.ModuleType("win32com"))
    monkeypatch.setitem(sys.modules, "win32com.client", None)


def test_pdf_has_stale_toc_detects_placeholder(tmp_path):
    stale = tmp_path / "stale.pdf"
    _write_pdf(stale, TOC_PLACEHOLDER)
    assert compiler._pdf_has_stale_toc(str(stale)) is True

    clean = tmp_path / "clean.pdf"
    _write_pdf(clean, "关于××工作的通知")
    assert compiler._pdf_has_stale_toc(str(clean)) is False


def test_pdf_has_stale_toc_tolerates_broken_file(tmp_path):
    """读不开的 PDF 一律按"无法判定"返回 False（不得影响主流程）。"""
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 broken")
    assert compiler._pdf_has_stale_toc(str(broken)) is False


@pytest.mark.skipif(sys.platform != "win32", reason="COM 分支只在 Windows 生效")
def test_docx_to_pdf_updates_fields_before_save(tmp_path, monkeypatch):
    """P1-03 核心：SaveAs2 之前必须 Fields.Update() + Repaginate()。"""
    src = _docx_stub(tmp_path)
    out = tmp_path / "out.pdf"
    log: list = []
    monkeypatch.setattr(compiler.shutil, "which", lambda name: None)   # 跳过 LibreOffice
    _install_fake_com(monkeypatch, log, out)

    assert compiler.docx_to_pdf(str(src), str(out)) == str(out)
    log = [x for x in log if not x.startswith("Dispatch:") and x not in ("Open", "Close", "Quit")]
    assert log == ["Fields.Update", "Repaginate", "SaveAs2"], f"顺序不对：{log}"


@pytest.mark.skipif(sys.platform != "win32", reason="COM 分支只在 Windows 生效")
def test_docx_to_pdf_survives_field_update_failure(tmp_path, monkeypatch):
    """更新域失败只记日志，不得让整次导出失败（转换产物仍可用）。"""
    src = _docx_stub(tmp_path)
    out = tmp_path / "out2.pdf"
    log: list = []
    monkeypatch.setattr(compiler.shutil, "which", lambda name: None)
    _install_fake_com(monkeypatch, log, out, fail_update=True)

    assert compiler.docx_to_pdf(str(src), str(out)) == str(out)
    assert "Repaginate" in log and "SaveAs2" in log
    assert out.exists()


def test_libreoffice_stale_toc_is_not_returned_silently(tmp_path, monkeypatch):
    """LibreOffice 产出"目录仍是占位文字"的 PDF 时，不能直接当成功返回。"""
    src = _docx_stub(tmp_path)
    out = tmp_path / "lo.pdf"
    _block_com(monkeypatch)                       # 没有 COM 可兜底
    monkeypatch.setattr(compiler.shutil, "which", lambda name: "/fake/soffice")

    def _fake_run(cmd, **kwargs):
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        stem = Path(cmd[-1]).stem
        _write_pdf(outdir / f"{stem}.pdf", TOC_PLACEHOLDER)
        return types.SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(compiler.subprocess, "run", _fake_run)

    got = compiler.docx_to_pdf(str(src), str(out))
    assert got == str(out) and out.exists()
    # 退回给用户的仍是那份 PDF，但目录确实是占位文字（已在日志里告警）
    assert compiler._pdf_has_stale_toc(got) is True


def test_libreoffice_clean_toc_returns_immediately(tmp_path, monkeypatch):
    """LibreOffice 产出的 PDF 目录正常时直接返回，不再去启动 COM。"""
    src = _docx_stub(tmp_path)
    out = tmp_path / "lo_ok.pdf"
    monkeypatch.setattr(compiler.shutil, "which", lambda name: "/fake/soffice")
    called = {"dispatch": 0}

    def _fake_run(cmd, **kwargs):
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        stem = Path(cmd[-1]).stem
        _write_pdf(outdir / f"{stem}.pdf", "关于××工作的通知")
        return types.SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(compiler.subprocess, "run", _fake_run)

    pkg = types.ModuleType("win32com")
    client = types.ModuleType("win32com.client")

    def _dispatch(progid):
        called["dispatch"] += 1
        raise AssertionError("目录正常时不应再尝试 COM")

    client.Dispatch = _dispatch
    pkg.client = client
    monkeypatch.setitem(sys.modules, "win32com", pkg)
    monkeypatch.setitem(sys.modules, "win32com.client", client)

    assert compiler.docx_to_pdf(str(src), str(out)) == str(out)
    assert called["dispatch"] == 0


def test_generated_docx_carries_toc_field_and_update_fields(tmp_path):
    """验收口径：TOC 域存在（dirty）+ settings 含 updateFields（配合 P1-03）。"""
    docx_path = _docx_stub(tmp_path)
    with zipfile.ZipFile(docx_path) as zf:
        document = zf.read("word/document.xml").decode("utf-8")
        settings = zf.read("word/settings.xml").decode("utf-8")
    assert "TOC" in document and "instrText" in document
    assert 'w:dirty="true"' in document
    assert TOC_PLACEHOLDER in document
    assert "updateFields" in settings


# ================================================================ P1-23 双冒号
_RECIPIENT_KINDS = ["指示", "通知", "通报", "报告", "请示", "批复", "函"]


def _recipient_lines(text: str) -> list[str]:
    return [ln for ln in text.splitlines() if ln.strip()]


@pytest.mark.parametrize("kind", _RECIPIENT_KINDS)
def test_recipient_with_fullwidth_colon_no_double(kind):
    """用户输入自带全角冒号：不得出现"：："。"""
    sk = skeletons.get(kind)
    text = sk.render(title="关于××的通知", org="××单位", matter="××事项",
                     recipients="各单位：", date="2026年8月30日")
    assert "：：" not in text
    assert "各单位：" in text
    assert text.count("各单位：") == 1
    # 主送行是独立一行，且以单个冒号收尾
    assert "各单位：" in _recipient_lines(text)


@pytest.mark.parametrize("kind", _RECIPIENT_KINDS)
def test_recipient_with_halfwidth_colon_normalized(kind):
    """半角冒号在公文里不合规：归一为全角，且只有一个。"""
    sk = skeletons.get(kind)
    text = sk.render(title="关于××的通知", org="××单位", matter="××事项",
                     recipients="各单位:", date="2026年8月30日")
    assert "各单位:" not in text
    assert "各单位：" in text
    assert text.count("各单位：") == 1


@pytest.mark.parametrize("kind", _RECIPIENT_KINDS)
def test_recipient_without_colon_keeps_template_colon(kind):
    """用户没带冒号：模板的冒号照旧（行为不变）。"""
    sk = skeletons.get(kind)
    text = sk.render(title="关于××的通知", org="××单位", matter="××事项",
                     recipients="各单位", date="2026年8月30日")
    assert "各单位：" in text
    assert "：：" not in text


def test_recipient_multiple_trailing_colons_collapse_to_one():
    sk = skeletons.get("通知")
    text = sk.render(title="T", org="O", matter="M", recipients="各单位：：", date="D")
    assert "：：" not in text
    assert "各单位：" in text


def test_recipient_trailing_space_before_colon():
    sk = skeletons.get("通知")
    text = sk.render(title="T", org="O", matter="M", recipients="各单位 ：", date="D")
    assert "各单位：" in text
    assert "：：" not in text


def test_recipient_empty_keeps_single_colon():
    """空主送：只保留模板冒号，不得拼出"：："或空行。"""
    sk = skeletons.get("通知")
    text = sk.render(title="T", org="O", matter="M", recipients="", date="D")
    assert "：：" not in text
    assert "\n：\n" in text


def test_skeleton_without_recipient_untouched():
    """无主送机关的骨架（命令/通告/纪要）不受影响：不留孤零零的冒号行。"""
    for kind in ("命令", "通告", "纪要"):
        sk = skeletons.get(kind)
        assert sk.recipient == ""
        text = sk.render(title="T", org="O", matter="M", date="D")
        assert "\n：\n" not in text
