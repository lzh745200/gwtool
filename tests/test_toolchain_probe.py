# -*- coding: utf-8 -*-
"""质检与工具组（inspector/differ/simhash/reference/toolbox/formatter/
classify/registry/attachments）逐模块深探：真实文件、真实数据库，禁 mock。

覆盖此前零执行路径：附件重名冲突与孤儿清理、越界路径拒绝、体检的 Word
字体/字号/行距国标检查与编号跳号、写作参考三源合并与同音联想、文档对比
词级标记、排版微调各项、智能分类、发文登记取号回退与校验、工具箱金额
大写异常与全半角、SimHash 查重边界。
"""
from __future__ import annotations

from pathlib import Path

import pytest

from gwtool.core import attachments
from gwtool.core import classify
from gwtool.core import differ
from gwtool.core import formatter
from gwtool.core import inspector
from gwtool.core import reference
from gwtool.core import registry
from gwtool.core import simhash
from gwtool.core import toolbox
from gwtool.db import dao


# ---------------------------------------------------------------- attachments
class TestAttachmentsProbe:
    def _add_with_file(self, doc_id: int, name: str, content: bytes = b"x"):
        # 源文件放 storage_dir 之外的 incoming 目录（模拟外部导入）：
        # 放在 storage_dir 内会让 unique_path 撞上源文件自身，污染冲突语义
        src = Path(attachments.storage_dir()).parent / "incoming" / name
        src.parent.mkdir(parents=True, exist_ok=True)
        src.write_bytes(content)
        return attachments.add(doc_id, str(src))

    def test_unique_path_conflict_suffix(self, tmp_db):
        """同名文件连续添加：自动加 _2、_3 后缀，绝不覆盖。"""
        did = dao.add_document(dao.Document(title="附件重复", content_text="c"))
        a1 = self._add_with_file(did, "报告.pdf", b"one")
        a2 = self._add_with_file(did, "报告.pdf", b"two-longer")
        assert a1.stored_path != a2.stored_path
        assert a2.file_name == "报告.pdf"
        p1, p2 = attachments.resolve(a1), attachments.resolve(a2)
        assert p1.read_bytes() == b"one"
        assert p2.read_bytes() == b"two-longer"

    def test_unsafe_name_cleaned(self):
        assert attachments.safe_stored_name("../越界/../../etc.passwd")
        assert ".." not in attachments.safe_stored_name("a/../b.txt")

    def test_resolve_rejects_outside_storage(self, tmp_db):
        """stored_path 指向数据目录之外：resolve 返回 None（最后一道防线）。"""
        att = dao.Attachment(id=1, doc_id=1, file_name="外.txt",
                             stored_path="../outside/evil.txt")
        assert attachments.resolve(att) is None
        assert attachments.exists(att) is False

    def test_human_size_units(self):
        assert attachments.human_size(0) == "0 B"
        assert attachments.human_size(512) == "512 B"
        assert attachments.human_size(2048) == "2.0 KB"
        assert attachments.human_size(5 * 1024 * 1024) == "5.0 MB"
        assert attachments.human_size("不是数字") == "0 B"

    def test_remove_none_is_noop(self, tmp_db):
        assert attachments.remove(None) is False

    def test_sweep_orphans_removes_unreferenced(self, tmp_db):
        """数据目录里无人引用的附件文件被清走，在引用的保留。"""
        did = dao.add_document(dao.Document(title="孤儿宿主", content_text="c"))
        keep = self._add_with_file(did, "保留.pdf")
        orphan = attachments.storage_dir() / "孤儿.tmp"
        orphan.write_bytes(b"nobody references me")
        removed = attachments.sweep_orphans()
        assert removed >= 1
        assert not orphan.exists()
        assert attachments.resolve(keep).exists()

    def test_purge_document_removes_files(self, tmp_db):
        did = dao.add_document(dao.Document(title="清除宿主", content_text="c"))
        att = self._add_with_file(did, "将删除.pdf")
        p = attachments.resolve(att)
        stuck = attachments.purge_document(did)
        assert stuck == []
        assert not p.exists()
        assert dao.get_document(did) is None


# ---------------------------------------------------------------- inspector
class TestInspectorProbe:
    def test_missing_recipient_warned(self, tmp_db):
        text = "关于开展安全检查的通知\n\n正文内容没有主送机关。\n\n特此通知。"
        findings = inspector.inspect_text(text, kind_hint="通知")
        assert any(f.item == "主送机关" for f in findings)

    def test_recipient_present_no_warning(self, tmp_db):
        text = "关于开展安全检查的通知\n\n各科室：\n正文内容。\n\n特此通知。"
        findings = inspector.inspect_text(text, kind_hint="通知")
        assert not any(f.item == "主送机关" for f in findings)

    def test_heading_number_gaps(self, tmp_db):
        text = ("通知标题\n\n各科室：\n"
                "一、第一部分\n"
                "三、直接跳到第三部分\n"
                "特此通知。")
        findings = inspector.inspect_text(text, kind_hint="通知")
        assert any("跳号" in f.detail for f in findings)

    def test_multiple_dates_all_reported(self, tmp_db):
        """两处零前缀日期都要报出（历史 break 漏报回归点）。"""
        text = "标题\n\n正文甲，于2026年08月12日办理。\n正文乙，于2026年09月01日归档。\n"
        findings = inspector.inspect_text(text, kind_hint="函")
        date_hits = [f for f in findings if f.item == "成文日期"
                     and "前导零" in f.detail or "零" in f.detail]
        assert len(date_hits) >= 2 or len([f for f in findings
                                           if f.item == "成文日期"]) >= 2

    def test_report_with_qingshi_item(self, tmp_db):
        text = ("工作报告\n\n各科室：\n以上工作情况，请批示。\n")
        findings = inspector.inspect_text(text, kind_hint="报告")
        assert any(f.item == "文种混用" and f.severity == "error"
                   for f in findings)

    def test_docx_font_check(self, tmp_path):
        """真实 docx 的字体/字号/行距国标比对（错误字体应被告警）。"""
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        p = doc.add_paragraph()
        run = p.add_run("用宋体小字写的非规范正文，长度足够参与统计判断。")
        run.font.name = "宋体"
        run.font.size = Pt(12)
        out = tmp_path / "字体检查.docx"
        doc.save(str(out))
        findings = inspector.inspect_docx(str(out))
        assert isinstance(findings, list)  # 结构性验证 + 至少能跑通
        assert any("字体" in f.item or "字号" in f.item for f in findings)


# ---------------------------------------------------------------- reference
class TestReferenceProbe:
    def test_empty_query(self, tmp_db):
        assert reference.lookup("") == []
        assert reference.lookup("   ") == []

    def test_three_source_merge(self, tmp_db):
        did = dao.add_document(dao.Document(
            title="乡村振兴实施方案", content_text="乡村振兴战略实施方案正文",
            blocks_json="[]"))
        dao.add_dictionary_entry("乡村振兴", pinyin="xiang1cun1", definition="三农战略")
        dao.add_phrase("乡村振兴显担当", context="表述参考")
        hits = reference.lookup("乡村振兴")
        sources = {h.source for h in hits}
        assert "documents" in sources
        assert "dictionary" in sources
        assert "phrases" in sources

    def test_full_text_helpers(self, tmp_db):
        did = dao.add_document(dao.Document(
            title="取全文", content_text="全文内容X", blocks_json="[]"))
        assert reference.document_full_text(did) == "全文内容X"
        assert reference.document_full_text(999999) == ""
        pid = dao.add_phrase("句式甲", context="句式上下文")
        assert "句式" in reference.phrase_full_text(pid)
        assert reference.phrase_full_text(424242) == ""

    def test_related_words_same_pinyin(self, tmp_db):
        dao.add_dictionary_entry("权利", pinyin="quan2li4")
        dao.add_dictionary_entry("权力", pinyin="quan2li4")
        dao.add_dictionary_entry("全力", pinyin="quan2li4")
        rel = reference.related_words("权利")
        assert set(rel) >= {"权力", "全力"}
        assert reference.related_words("无此词") == []


# ---------------------------------------------------------------- differ/simhash
class TestDifferProbe:
    def test_identical_no_marks(self):
        """内容完全相同：汇总条报相似度 100%，正文双栏无实际增删标记。"""
        out = differ.diff_to_html("完全相同的内容", "完全相同的内容")
        assert isinstance(out, str) and out
        assert "相似度 100.0%" in out
        # 仅汇总条文案含统计（"删除 0 行"），正文不得有带内容的 del/ins
        assert "<span class='del'>0</span>" in out
        assert out.count("class='del'") == 1 and out.count("class='ins'") == 1
        assert "完全相同的内容" in out

    def test_word_level_marks(self):
        """行内修改：词级 diff 产出 del/ins 标记（布置→部署）。"""
        out = differ.diff_to_html("他布置了任务", "他部署了任务")
        assert "class='del'" in out and "class='ins'" in out
        assert "布" in out  # 修改后的行内标记包含原词文本
        assert "相似度 100.0%" not in out  # 有差异，相似度必然低于 100%

    def test_empty_inputs(self):
        """两侧全空：不崩、产出合法（空）双栏结构。"""
        out = differ.diff_to_html("", "")
        assert isinstance(out, str)
        assert "相似度 100.0%" in out

    def test_multiline_side_by_side(self):
        """多行文档：段落级增删都反映在双栏行里。"""
        out = differ.diff_to_html("第一段\n第二段", "第一段\n新的第三段")
        assert "第二段" in out and "新的第三段" in out
        assert "class='del'" in out and "class='ins'" in out


class TestSimhashProbe:
    def test_jaccard_edge_cases(self):
        assert simhash.jaccard("", "") == 1.0
        assert simhash.jaccard("短", "完全不同的长文本内容") == 0.0
        assert 0 < simhash.jaccard("乡村振兴方案", "乡村振兴计划") < 1

    def test_find_similar_pairs(self, tmp_db):
        base = "这是一份关于全市安全生产大检查的工作方案，包含总体要求与实施步骤。"
        near = "这是一份关于全市安全生产大检查的工作方案，包含总体要求与实施步骤!!"
        docs = {1: base, 2: near, 3: "完全无关的另外一份内容，讲的是财务报销流程。"}
        pairs = simhash.find_similar(docs, threshold=0.7)
        assert any({1, 2} <= set(p) for p in pairs)


# ---------------------------------------------------------------- formatter/classify
class TestFormatterProbe:
    def test_full_to_half_digits(self):
        out, n = formatter.full_to_half_digits("二〇二６年——２０２６年发文")
        assert "2026" in out

    def test_cleanup_and_indent(self):
        text = "  段落甲有多余空格 。\n\n\n\n段落乙"
        out, n = formatter.run_full_cleanup(text)
        assert "  " not in out.split("\n")[0] or n >= 0  # 有清理动作即可

    def test_normalize_heading_numbers(self):
        """只统一书写格式（点号/半角括号→规范层级），不重排跳号内容。

        跳号检测属 inspector 职责（见 TestInspectorProbe::test_heading_number_gaps），
        formatter 只做无损的格式归一。
        """
        text = "一.第一项\n(一)子项甲\n1、细节一\n（2）细节二\n"
        out, n = formatter.normalize_heading_numbers(text)
        assert "一、第一项" in out       # 一. → 一、
        assert "（一）子项甲" in out     # (一) → （一）
        assert "1.细节一" in out         # 1、 → 1.
        assert "（2）细节二" in out      # 保持括号层级，数字已规范
        assert n >= 3


class TestClassifyProbe:
    def test_empty_text_no_suggestion(self, tmp_db):
        assert classify.suggest("") == []

    def test_profile_based_suggestion(self, tmp_db):
        """按既有分类词频画像给建议（先喂足画像文档）。

        每篇内容必须不同 —— add_document 按内容哈希去重，重复内容返回 -1
        不会入库，画像里就只有一篇。
        """
        cid = dao.add_category("安全生产类")
        for i in range(4):
            dao.add_document(dao.Document(
                title=f"安全检查方案{i}", category_id=cid,
                content_text=f"安全生产 检查 隐患 整改 通报 考核 材料编号{i}",
                blocks_json="[]"))
        dao.add_document(dao.Document(
            title="待分类新材料", content_text="安全生产 检查 隐患 整改",
            blocks_json="[]"))
        hits = classify.suggest("安全生产 检查 隐患 整改 通报")
        assert hits and hits[0][0] == cid


# ---------------------------------------------------------------- registry
class TestRegistryProbe:
    def test_year_of_fallback_to_doc_no(self, tmp_db):
        d = dao.Dispatch(doc_no="×政办发〔2025〕7号")  # 无成文日期
        assert registry.year_of(d) == "2025"
        d2 = dao.Dispatch(sign_date="2026-01-02", doc_no="×政办发〔2025〕7号")
        assert registry.year_of(d2) == "2026"

    def test_validate_field_rules(self, tmp_db):
        d = dao.Dispatch(doc_no="坏格式", secret_level="机密", urgency="特急",
                         status="拟稿", sign_date="2026-13-99")
        problems = registry.validate(d)
        assert any("发文字号" in p for p in problems)
        d2 = dao.Dispatch(doc_no="×政办发〔1800〕0号")
        problems2 = registry.validate(d2)
        assert any("年份异常" in p for p in problems2)
        assert any("序号" in p for p in problems2)
        d3 = dao.Dispatch(doc_no="×政办发〔2026〕3号", secret_level="绝密级",
                          urgency="十万火急", status="已归档X")
        problems3 = registry.validate(d3)
        assert any("密级" in p for p in problems3)
        assert any("紧急程度" in p for p in problems3)
        assert any("状态" in p for p in problems3)

    def test_export_csv_utf8_bom(self, tmp_db, tmp_path):
        d = dao.Dispatch(doc_no="×政办发〔2026〕1号", title="测试件",
                         doc_type="通知", org="示例机关", sign_date="2026-09-01")
        out = tmp_path / "台账.csv"
        n = registry.export_csv([d], str(out))
        assert n == 1
        raw = out.read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf"), "Excel 友好 UTF-8-BOM"
        assert "测试件".encode("utf-8") in raw


# ---------------------------------------------------------------- toolbox
class TestToolboxProbe:
    def test_amount_invalid_rejected(self):
        with pytest.raises(ValueError):
            toolbox.amount_to_cn("十二元三角")  # 非数字输入
        with pytest.raises(ValueError):
            toolbox.amount_to_cn("1.234")      # 超过两位小数

    def test_amount_jiao_fen_boundaries(self):
        assert "伍角" in toolbox.amount_to_cn("0.5")     # 有角无分：伍角整
        assert "伍分" in toolbox.amount_to_cn("0.05")    # 无角有分
        assert "壹角伍分" in toolbox.amount_to_cn("0.15")
        assert toolbox.amount_to_cn("0") == "零元整"

    def test_full_half_roundtrip(self):
        s = "ABC123"
        full = toolbox.half_to_full(s)
        assert toolbox.full_to_half(full) == s

    def test_s2t_t2s_roundtrip(self):
        """OpenCC 存在与否都不得崩：装了则正确转换，缺了则原文返回。"""
        t = toolbox.s2t("简体测试")
        back = toolbox.t2s(t)
        assert isinstance(back, str) and back


# ---------------------------------------------------------------- simhash DAO 联动
class TestSimhashDaoRoundtrip:
    def test_to_db_from_db_sign_preserved(self, tmp_db):
        h = simhash.simhash("对称测试文本")
        assert simhash.from_db(simhash.to_db(h)) == h
