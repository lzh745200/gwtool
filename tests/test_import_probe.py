# -*- coding: utf-8 -*-
"""导入解析链（importer + parsers）逐模块深探：全部真实样张文件，禁 mock。

真实样张在测试内动态生成（python-docx / pymupdf / 原始字节），覆盖：
  - parse_any 全格式分发（docx/doc/txt/rtf/md/html/pdf/不支持格式）；
  - 损坏文件与空文件的异常路径（单文件失败不影响批处理）；
  - 扫描版 PDF 无 OCR 时的用户提示；OCR 文本 -> DocTree 转换；
  - batch_import 进度回调；txt 多编码探测边界；
  - .doc 真机解析：WPS COM 把 docx 另存为 .doc 后走四级降级链主路径
    （无 WPS/Word 的环境自动跳过）。
"""
from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from gwtool.core.importer import batch_import, parse_any
from gwtool.core.model import HEADING, PARAGRAPH, TABLE


# ---------------------------------------------------------------- 样张工厂
def _mk_docx(path: Path, paras: list | None = None) -> Path:
    from docx import Document
    doc = Document()
    doc.add_heading("样张标题", level=1)
    for p in (paras or ["第一段正文。", "第二段正文。"]):
        doc.add_paragraph(p)
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "表格甲"
    doc.save(str(path))
    return path


def _mk_pdf_text(path: Path) -> Path:
    import pymupdf as fitz
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 90), "文字版PDF标题", fontname="china-s", fontsize=16)
    page.insert_text((72, 120), "这是文字版PDF的正文内容。", fontname="china-s",
                     fontsize=12)
    doc.save(str(path))
    doc.close()
    return path


def _mk_pdf_scan(path: Path) -> Path:
    """纯图页 PDF（无文本层），模拟扫描件。"""
    import pymupdf as fitz
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.draw_rect(fitz.Rect(50, 50, 500, 300), color=(0, 0, 0))
    doc.save(str(path))
    doc.close()
    return path


# ---------------------------------------------------------------- 分发与异常
class TestParseAnyDispatch:
    def test_docx(self, tmp_path):
        r = parse_any(str(_mk_docx(tmp_path / "a.docx")))
        assert r.ok and r.tree.blocks
        assert any(b.type == TABLE for b in r.tree.blocks)

    def test_txt(self, tmp_path):
        p = tmp_path / "b.txt"
        p.write_text("会议通知\n关于召开年度会议的通知正文。", encoding="utf-8")
        r = parse_any(str(p))
        assert r.ok and r.tree.title

    def test_md(self, tmp_path):
        p = tmp_path / "c.md"
        p.write_text("# 一级标题\n\n正文段落。\n\n- 列表项一\n- 列表项二\n",
                     encoding="utf-8")
        r = parse_any(str(p))
        assert r.ok and r.tree.blocks

    def test_html(self, tmp_path):
        p = tmp_path / "d.html"
        p.write_text("<html><body><h1>网页标题</h1><p>网页正文。</p>"
                     "</body></html>", encoding="utf-8")
        r = parse_any(str(p))
        assert r.ok and any(b.text == "网页标题" for b in r.tree.blocks)

    def test_rtf(self, tmp_path):
        p = tmp_path / "e.rtf"
        # 中文用 RTF 标准 \uN 转义（'富'=U+5BCC，'文'=U+6587）
        p.write_text(
            r"{\rtf1\ansi RTF \b rich text\b0 \u23500? \u25991? test.}",
            encoding="ascii")
        r = parse_any(str(p))
        assert r.ok, r.error
        assert "富" in "".join(b.text for b in r.tree.blocks)

    def test_pdf_text_layer(self, tmp_path):
        r = parse_any(str(_mk_pdf_text(tmp_path / "f.pdf")))
        assert r.ok, r.error
        assert r.tree.blocks

    def test_scan_pdf_without_ocr_hint(self, tmp_path):
        """无文本层 PDF 且本机无 OCR：给出可执行的安装指引而非静默失败。"""
        from gwtool.core.ocr import available as ocr_av
        r = parse_any(str(_mk_pdf_scan(tmp_path / "scan.pdf")))
        if ocr_av():
            pytest.skip("本机已装 OCR，走的是 OCR 兜底路径")
        assert not r.ok
        assert "Tesseract" in r.error

    def test_unsupported_ext(self, tmp_path):
        p = tmp_path / "g.xyz"
        p.write_text("任意内容", encoding="utf-8")
        r = parse_any(str(p))
        assert not r.ok and "不支持的格式" in r.error

    def test_corrupted_docx_reported_not_raised(self, tmp_path):
        """损坏 docx：返回错误结果而非抛异常（批处理健壮性契约）。"""
        p = tmp_path / "broken.docx"
        p.write_bytes(b"PK\x03\x04" + b"\x00" * 256)  # zip 魔数 + 垃圾
        r = parse_any(str(p))
        assert not r.ok and r.error

    def test_empty_txt(self, tmp_path):
        p = tmp_path / "empty.txt"
        p.write_text("", encoding="utf-8")
        r = parse_any(str(p))
        assert not r.ok and "未提取到文字" in r.error

    def test_missing_file(self, tmp_path):
        r = parse_any(str(tmp_path / "不存在.docx"))
        assert not r.ok and r.error


class TestTxtEncodings:
    @pytest.mark.parametrize("enc", ["utf-8", "gbk"])
    def test_common_encodings(self, tmp_path, enc):
        p = tmp_path / f"enc_{enc}.txt"
        text = "公文标题行\n这是正文内容，用于编码探测验证。"
        p.write_text(text, encoding=enc)
        r = parse_any(str(p))
        assert r.ok, f"{enc} 解析失败：{r.error}"
        assert "正文内容" in "".join(b.text for b in r.tree.blocks)

    def test_big5_traditional(self, tmp_path):
        """BIG5 是繁体编码：样张必须用繁体字（简体字根本无法编码为 BIG5）。"""
        p = tmp_path / "enc_big5.txt"
        text = "公文標題行\n這是正文內容，用於編碼探測驗證。"
        p.write_text(text, encoding="big5")
        r = parse_any(str(p))
        assert r.ok, f"BIG5 解析失败：{r.error}"
        assert "正文內容" in "".join(b.text for b in r.tree.blocks)

    def test_utf8_bom(self, tmp_path):
        p = tmp_path / "bom.txt"
        p.write_text("带BOM的文本\n正文。", encoding="utf-8-sig")
        r = parse_any(str(p))
        assert r.ok
        joined = "".join(b.text for b in r.tree.blocks)
        assert "带BOM" in joined


class TestOcrTextToTree:
    """OCR 文本 -> DocTree 的结构化（标题提取/标题行/段落）。"""

    def test_structure(self):
        from gwtool.core.importer import _text_to_tree
        text = ("关于切实做好秋季安全工作的通知\n"
                "各科室、各单位：\n"
                "一、提高思想认识\n"
                "这是正文段落，内容较长所以不会被判为标题行。\n"
                "二、强化责任落实\n"
                "落款单位名称也应当作正文处理。")
        tree = _text_to_tree(text, title_hint="scan.png")
        assert tree.title == "关于切实做好秋季安全工作的通知"
        kinds = [(b.type, b.text) for b in tree.blocks]
        assert kinds[0][0] == HEADING
        assert any(t == HEADING and "一、" in s for t, s in kinds)
        assert any(t == PARAGRAPH and "正文段落" in s for t, s in kinds)

    def test_empty_text(self):
        from gwtool.core.importer import _text_to_tree
        tree = _text_to_tree("")
        assert tree.blocks == []


class TestBatchImport:
    def test_progress_callback_and_mixed_results(self, tmp_path):
        good = str(_mk_docx(tmp_path / "ok.docx"))
        bad = str(tmp_path / "bad.xyz")
        bad_content = Path(bad)
        bad_content.write_text("x", encoding="utf-8")
        seen: list[tuple[int, int, str]] = []
        results = batch_import([good, bad],
                               progress_cb=lambda i, t, p: seen.append((i, t, p)))
        assert len(results) == 2
        assert results[0].ok and not results[1].ok
        assert [s[0] for s in seen] == [1, 2]
        assert seen[1][2] == bad


# ---------------------------------------------------------------- .doc 真机
class TestDocRealConversion:
    @pytest.fixture(scope="class")
    def real_doc(self, tmp_path_factory):
        """用本机 WPS/Word COM 把 docx 另存为 .doc（97-2003）。

        这是四级降级链的主路径（COM 保真最高）。无 COM 组件则整组跳过。
        """
        if shutil.os.name != "nt":
            pytest.skip("COM 仅 Windows")
        try:
            import win32com.client
        except ImportError:
            pytest.skip("未安装 pywin32")
        src = _mk_docx(tmp_path_factory.mktemp("docprobe") / "src.docx")
        dst = src.with_suffix(".doc")
        app = None
        try:
            for progid in ("kwps.Application", "wps.Application", "Word.Application"):
                try:
                    app = win32com.client.Dispatch(progid)
                    break
                except Exception:
                    continue
            if app is None:
                pytest.skip("本机无 WPS/Word COM")
            app.Visible = False
            doc = app.Documents.Open(str(src.resolve()), ReadOnly=True)
            doc.SaveAs2(str(dst.resolve()), FileFormat=0)  # wdFormatDocument
            doc.Close(False)
        finally:
            if app is not None:
                app.Quit()
        assert dst.exists() and dst.stat().st_size > 0
        return dst

    def test_parse_doc_via_com(self, real_doc):
        r = parse_any(str(real_doc))
        assert r.ok, f"真实 .doc 解析失败：{r.error}"
        joined = "".join(b.text for b in r.tree.blocks)
        assert "第一段正文。" in joined

    def test_wps_ooxml_sniffing(self, tmp_path):
        """.wps 扩展名 + OOXML 内容：按魔数路由到 docx 解析器。"""
        p = tmp_path / "ooxml.wps"
        _mk_docx(p)  # zip 魔数 PK\x03\x04
        r = parse_any(str(p))
        assert r.ok, r.error
        assert "第一段正文。" in "".join(b.text for b in r.tree.blocks)

    def test_garbage_wps_graceful(self, tmp_path):
        """.wps 扩展名 + 非魔数垃圾：走 .doc 链后优雅失败，不崩进程。"""
        p = tmp_path / "junk.wps"
        p.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 512)  # OLE 魔数 + 空
        r = parse_any(str(p))
        assert not r.ok and r.error  # 错误信息返回，绝不抛 SEH
